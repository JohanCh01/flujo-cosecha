"""
Temperaturas de los carros (traslado de campo a planta) para la página carros.html del Flujo de cosecha.

Corre en TU LAPTOP. Lee los Excel de los sensores (hoja «Detalles») de dos carpetas:
    <carpeta>\\Carro Metal     dos sensores: uno del CHOFER (fondo) y uno de la PUERTA
    <carpeta>\\Carro madera    igual
De cada Excel saca:
  - si el sensor es del CHOFER o de la PUERTA: lo dice el título «SENSOR # … (FONDO - CHOFER)» o «(PUERTA)»
    que está arriba a la derecha (I4, J3, I5…; se busca en esa zona)
  - las lecturas: Hora (columna B) y Temperatura (columna C)
  - el tramo que se analiza:
      inicio = la fila pintada de AMARILLO   · fin = la fila pintada de ROJO (llegada a planta), las dos incluidas
      Carro de METAL: el inicio NO es el amarillo, sino la primera lectura desde las 12:00 del día INICIO_METAL
      (nunca antes de las 12:00).
Lo guarda cifrado en carros.json (misma lista de DNIs y clave de administrador que el resto) y lo sube a GitHub.

Uso (en Git Bash, dentro de la carpeta flujo-cosecha):
    python carros_temperatura.py --carpeta "C:/Users/jchinguel/Desktop/CARROS"
    python carros_temperatura.py --carpeta "…" --inicio-metal "2026-10-28 12:00"
    python carros_temperatura.py --carpeta "…" --sin-subir        (solo muestra el resumen)
"""

import argparse
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
import recepcion_sql as rs          # misma clave de administrador, mismo cifrado y misma subida a GitHub
af, log, git = rs.af, rs.log, rs.git

INICIO_METAL = "2026-10-28 12:00"   # metal: primera lectura desde esta hora (nunca antes)
AMARILLO, ROJO = "FFFFC000", "FFFF0000"
RUTA = AQUI / "carros.json"


def norm(t):
    t = unicodedata.normalize("NFD", str(t or ""))
    return "".join(c for c in t if unicodedata.category(c) != "Mn").upper()


def fecha(v):
    if isinstance(v, datetime):
        return v
    s = str(v or "").strip()
    for f in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return None


def numero(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def color(celda):
    f = celda.fill
    return (f.fgColor.rgb or "").upper() if f is not None and f.fill_type and isinstance(f.fgColor.rgb, str) else ""


def leer_sensor(ruta):
    import openpyxl
    ws = openpyxl.load_workbook(ruta)["Detalles"]
    titulo, pos = None, None
    for fila in ws.iter_rows(min_row=1, max_row=20, min_col=5, max_col=20):
        for c in fila:
            if isinstance(c.value, str) and "SENSOR" in norm(c.value):
                titulo = " ".join(c.value.split())
                pos = "chofer" if "CHOFER" in norm(c.value) else "puerta" if "PUERTA" in norm(c.value) else None
                break
        if titulo:
            break
    if not pos:
        raise ValueError(f"{ruta.name}: no encontré el título «SENSOR # … (CHOFER / PUERTA)» en la hoja Detalles")
    lect, ini, fin = [], None, None
    for r in range(2, ws.max_row + 1):
        h, t = fecha(ws.cell(r, 2).value), numero(ws.cell(r, 3).value)
        if h is None or t is None:
            continue
        col = color(ws.cell(r, 2))
        if col == AMARILLO and ini is None:
            ini = len(lect)
        if col == ROJO:
            fin = len(lect)
        lect.append((h, t))
    if fin is None:
        raise ValueError(f"{ruta.name}: no encontré la fila en ROJO (llegada a planta)")
    return {"archivo": ruta.name, "titulo": titulo, "pos": pos, "lect": lect, "ini": ini, "fin": fin}


def carpeta_de(raiz, palabra):
    for d in sorted(raiz.iterdir()):
        if d.is_dir() and palabra in norm(d.name):
            return d
    sys.exit(f"No encontré dentro de {raiz} una carpeta con «{palabra.title()}» en el nombre.")


def main():
    ap = argparse.ArgumentParser(description="Temperaturas de los carros de metal y madera → carros.html")
    ap.add_argument("--carpeta", required=True, help="carpeta que tiene las carpetas «Carro Metal» y «Carro madera»")
    ap.add_argument("--inicio-metal", default=INICIO_METAL, help="metal: primera lectura desde esta hora (AAAA-MM-DD HH:MM)")
    ap.add_argument("--sin-subir", action="store_true")
    a = ap.parse_args()
    raiz = Path(a.carpeta)
    if not raiz.is_dir():
        sys.exit(f"No encuentro la carpeta {raiz}")
    desde_metal = fecha(a.inicio_metal)
    if not desde_metal:
        sys.exit("--inicio-metal debe ser como 2026-10-28 12:00")

    carros = {}
    for clave, palabra in (("metal", "METAL"), ("madera", "MADERA")):
        sens = {}
        for f in sorted(carpeta_de(raiz, palabra).glob("*.xls*")):
            if f.name.startswith("~$"):
                continue
            s = leer_sensor(f)
            if clave == "metal":                      # metal: desde las 12:00, nunca antes
                s["ini"] = next((i for i, (h, _) in enumerate(s["lect"]) if h >= desde_metal), None)
            if s["ini"] is None:
                sys.exit(f"{f.name}: no encontré el inicio del tramo")
            if s["pos"] in sens:
                sys.exit(f"Carro {clave}: hay dos sensores de {s['pos'].upper()} ({sens[s['pos']]['archivo']} y {f.name})")
            sens[s["pos"]] = s
            tramo = s["lect"][s["ini"]:s["fin"] + 1]
            ts = [t for _, t in tramo]
            log(f"{clave.upper():6} · {s['pos']:6} · {f.name}: {tramo[0][0]:%d/%m %H:%M} → {tramo[-1][0]:%d/%m %H:%M} · "
                f"{len(tramo)} lecturas · mín {min(ts)} · máx {max(ts)} · prom {sum(ts) / len(ts):.1f} °C")
        if set(sens) != {"chofer", "puerta"}:
            sys.exit(f"Carro {clave}: faltan sensores (encontré: {', '.join(sens) or 'ninguno'})")
        carros[clave] = {p: {"archivo": s["archivo"], "titulo": s["titulo"], "ini": s["ini"], "fin": s["fin"],
                             "h": [h.strftime("%Y-%m-%d %H:%M") for h, _ in s["lect"]], "t": [t for _, t in s["lect"]]}
                         for p, s in sens.items()}
    paquete = {"v": 1, "gen": datetime.now().strftime("%Y-%m-%dT%H:%M"), "inicioMetal": desde_metal.strftime("%Y-%m-%d %H:%M"), "carros": carros}
    if a.sin_subir:
        log("Prueba (--sin-subir): no se subió nada.")
        return

    clave = rs.clave_admin()
    for _ in range(3):
        git("pull", "--ff-only", "-q")
        dnis, it, sal = af.accesos(clave)
        af.RUTA_FITO = RUTA                            # mismo cifrado que fito.json, en su propio archivo
        af.guardar_fito(paquete, dnis, clave, it, sal)
        git("add", RUTA.name)
        git("commit", "-q", "-m", f"Temperaturas de carros ({datetime.now():%d/%m/%Y %H:%M})")
        try:
            git("push", "-q")
            log("Listo: subido. En 1–2 minutos se ve en la página «Temperaturas en carros» (Ctrl + F5).")
            return
        except RuntimeError:
            log("GitHub tenía un cambio más nuevo; vuelvo a intentarlo…")
            git("reset", "-q", "--keep", "HEAD~1")
    sys.exit("No se pudo subir después de 3 intentos. Prueba otra vez en unos minutos.")


if __name__ == "__main__":
    main()
