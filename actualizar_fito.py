"""
Actualización automática de la página fitosanitario.html (cruce cosecha vs. evaluación fitosanitaria).

Corre en GitHub Actions justo DESPUÉS de actualizar.py, dentro del mismo trabajo. Qué hace:

1. RECEPCIÓN: toma el mismo Excel de "Recepción de Materia Prima" que actualizar.py acaba de
   bajar de la intranet (no entra otra vez a la intranet) y se queda solo con lo que usa el cruce:
   fecha de Hora Cosecha, Unidad Agrícola, Proveedor, Campo, CLP, Categoría y Peso Planta,
   únicamente de las filas con Proceso = FR. Esos días reemplazan a los que ya estaban guardados.
2. FOSFINA: baja de Google Drive la base "BD_FOSFINA" completa (todo el año) con la cuenta de
   servicio, igual que bd.py, y REEMPLAZA la anterior.
2b. CALIDADES Y CALIBRES y LONGITUDES (para reporte.html): las baja de la misma carpeta de Drive y
   guarda solo los totales por día y unidad agrícola (PROCESO = FRESCO, STATUS = EVALUADO;
   CARACTERISTICA = NORMAL en calidades y no vacía en longitudes).
3. Guarda las dos bases juntas en fito.json, cifrado con la misma lista de DNIs y la misma clave
   de administrador que datos.json. La página fitosanitario.html lo abre con la sesión ya iniciada
   en el Flujo de cosecha.

Si algo falla aquí, el Flujo de cosecha no se ve afectado: datos.json ya quedó actualizado antes.

También corre en una PC, para cargar historia o probar:

    python actualizar_fito.py --recepcion Recepcion.xlsx --todo --fosfina BD_FOSFINA.xlsx

Credenciales (nunca van escritas aquí; en GitHub son "secrets"):
    ADMIN_CLAVE          la misma del Flujo de cosecha
    GOOGLE_CREDENCIALES  el contenido completo de credenciales_google.json (cuenta de servicio)
    DRIVE_CARPETA_ID     el id de la carpeta de Drive donde está la base de fosfina
"""

import argparse
import base64
import gzip
import hashlib
import io
import json
import math
import os
import re
import sys
import unicodedata
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

LIMA = timezone(timedelta(hours=-5))
RUTA_DATOS = Path(os.environ.get("DATOS_JSON", "datos.json"))   # de aquí salen la lista de DNIs y la sal
RUTA_FITO = Path(os.environ.get("FITO_JSON", "fito.json"))
NOMBRE_FOSFINA = os.environ.get("FOSFINA_NOMBRE", "BD_FOSFINA")  # parte del nombre del archivo en Drive


def log(msg):
    print(f"[{datetime.now(LIMA).strftime('%H:%M:%S')}] {msg}", flush=True)


class Detener(Exception):
    """Error esperado: se muestra el mensaje y no se publica nada."""


# ============================================================
# 1. LECTURA DE LOS EXCEL (mismas reglas que la página web)
# ============================================================

def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def _limpio(v):
    """Texto en mayúsculas, sin espacios repetidos ni al borde."""
    return " ".join(_txt(v).split()).upper()


def _clave_col(v):
    """Encabezado para comparar: sin tildes, en mayúsculas y sin espacios."""
    s = unicodedata.normalize("NFD", _txt(v))
    s = "".join(c for c in s if not ("̀" <= c <= "ͯ"))
    return re.sub(r"\s+", "", s.upper())


def _ymd(v):
    """Solo la fecha, como número AAAAMMDD (0 si no es una fecha). Si trae hora, se descarta."""
    def ok(y, m, d):
        return y * 10000 + m * 100 + d if 2000 <= y <= 2100 and 1 <= m <= 12 and 1 <= d <= 31 else 0
    if v is None or v == "":
        return 0
    if isinstance(v, datetime):
        return ok(v.year, v.month, v.day)
    if isinstance(v, date):
        return ok(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            d = date(1899, 12, 30) + timedelta(days=int(math.floor(v)))
        except (OverflowError, ValueError):
            return 0
        return ok(d.year, d.month, d.day)
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return ok(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)      # día/mes/año, con o sin hora detrás
    if m:
        y = int(m.group(3))
        return ok(y + 2000 if y < 100 else y, int(m.group(2)), int(m.group(1)))
    return 0


def _num(v):
    if isinstance(v, bool) or v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else 0.0
    s = str(v).strip()
    if not s:
        return 0.0
    pct = "%" in s
    m = re.match(r"^[+-]?(\d+\.?\d*|\.\d+)", s.replace("%", "").replace(",", "."))
    if not m:
        return 0.0
    n = float(m.group(0))
    return n / 100 if pct else n


def _redondeo(x):
    return int(math.floor(x + 0.5))


def _semana_iso(ymd):
    return date(ymd // 10000, ymd // 100 % 100, ymd % 100).isocalendar()[1]


def _dias_entre(base, ymd):
    return (date(ymd // 10000, ymd // 100 % 100, ymd % 100) - date(base // 10000, base // 100 % 100, base % 100)).days


def _sumar_dias(base, n):
    d = date(base // 10000, base // 100 % 100, base % 100) + timedelta(days=n)
    return d.year * 10000 + d.month * 100 + d.day


def _hojas(ruta, preferidas=()):
    """Devuelve (nombre, filas) de cada hoja; primero las de nombre preferido."""
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)
    nombres = sorted(wb.sheetnames, key=lambda n: 0 if _clave_col(n) in preferidas else 1)
    for n in nombres:
        yield n, [list(f) for f in wb[n].iter_rows(values_only=True)]
    wb.close()


def _encabezado(filas, cols, obligatorias):
    for r, fila in enumerate(filas[:20]):
        claves = [_clave_col(c) for c in fila]
        if all(k in claves for k in obligatorias):
            idx, faltan = {}, []
            for k, (nombres, etiqueta, req) in cols.items():
                at = next((claves.index(n) for n in nombres if n in claves), -1)
                idx[k] = at
                if at < 0 and req:
                    faltan.append(etiqueta)
            return r, idx, faltan
    return None


COLS_FOSFINA = {
    "fecha": (["FECHA"], "FECHA", True), "ua": (["UNIDADAGRICOLA"], "UNIDAD AGRICOLA", True),
    "prov": (["PROVEEDOR"], "PROVEEDOR", True), "turno": (["TURNO"], "TURNO", True),
    "ph": (["P.HELIOTHIS", "PHELIOTHIS"], "P.Heliothis", True), "ps": (["P.SPODOPTERA", "PSPODOPTERA"], "P.Spodoptera", True),
    "po": (["P.OTROS", "POTROS"], "P.Otros", True), "larva": (["LARVA", "LARVAS"], "Larva", True),
    "mod": (["MODULO"], "MODULO", False), "clase": (["CLASEDECAMPO"], "CLASE DE CAMPO", False),
    "status": (["STATUS", "ESTADO"], "STATUS", False), "sem": (["SEMANA"], "SEMANA", False),
    # para el Reporte diario (reporte.html); si alguna no está, esa tarjeta sale vacía y lo demás sigue igual
    "kilos": (["KILOS", "KG", "KILOS(KG)"], "KILOS", False),
    "color": (["COLOR", "SEMAFORO", "COLORELASMOPALPUS"], "COLOR", False),
    "pic": (["SUMADEPICADOS", "SUMAPICADOS", "PICADOS", "PICADOSELASMOPALPUS", "PICADOSELASMOSPALPUS", "PICADOSELASMO",
             "PICADOS(ELASMOPALPUS)", "PICADO", "%PICADOS", "%PICADOSELASMOPALPUS"], "Suma de picados (Elasmopalpus)", False),
    "prodi": (["PICADOSPRODIPLOSIS", "PICADOPRODIPLOSIS", "PRODIPLOSIS", "%PICADOSPRODIPLOSIS"], "Picados Prodiplosis", False),
    "proceso": (["PROCESO"], "PROCESO", False),     # el Reporte diario solo toma PROCESO = FRESCO
}

# Columnas que usa el Reporte diario; se avisa en el registro si falta alguna
COLS_REPORTE = ("kilos", "color", "pic", "prodi", "proceso")
IND_W = 7            # números por evaluación en el bloque 'ind'
PROCESO_REPORTE = "FRESCO"


def _num_o_nada(v):
    """Como _num, pero una celda vacía o que no es número devuelve None (Power BI la deja fuera del promedio)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    s = str(v).strip().replace("%", "").replace(",", ".")
    if not re.match(r"^[+-]?(\d+\.?\d*|\.\d+)$", s):
        return None
    return _num(v)


def leer_fosfina(ruta):
    """Base de evaluaciones fitosanitarias, en el formato compacto que lee la página."""
    error = "No encontré los encabezados FECHA y PROVEEDOR en la base de fosfina."
    for _, filas in _hojas(ruta, ("FOSFINA", "FITOSANITARIO")):
        enc = _encabezado(filas, COLS_FOSFINA, ["FECHA", "PROVEEDOR"])
        if not enc:
            continue
        r0, ix, faltan = enc
        if faltan:
            error = "Faltan columnas en la base de fosfina: " + ", ".join(faltan) + "."
            continue

        def celda(fila, k):
            return fila[ix[k]] if 0 <= ix[k] < len(fila) else None

        mapa, campos, tops, evs, estados, n_filas = {}, [], [], [], [], 0
        colores, ind = [], []        # datos extra del Reporte diario, uno por evaluación (mismo orden que evs)

        def millon(x):
            return -1 if x is None else _redondeo(x * 1e6)
        for fila in filas[r0 + 1:]:
            ua, prov, turno = _limpio(celda(fila, "ua")), _limpio(celda(fila, "prov")), _limpio(celda(fila, "turno"))
            if not ua and not prov and not turno:
                continue
            ymd = _ymd(celda(fila, "fecha"))
            if not ymd or not prov:
                continue
            n_filas += 1
            mod = _limpio(celda(fila, "mod"))
            if _clave_col(mod) == _clave_col(prov):       # en terceros el módulo repite al proveedor
                mod = ""
            llave = (ua, prov, mod, turno)
            fi = mapa.get(llave)
            if fi is None:
                fi = mapa[llave] = len(campos)
                campos.append([ua, "", prov, mod, turno])
                tops.append(0)
            if ymd >= tops[fi]:
                tops[fi] = ymd
                if ix["clase"] >= 0:
                    campos[fi][1] = _limpio(celda(fila, "clase"))
            st = _limpio(celda(fila, "status")) if ix["status"] >= 0 else ""
            if st not in estados:
                estados.append(st)
            sem = _redondeo(_num(celda(fila, "sem"))) if ix["sem"] >= 0 else 0
            if not 1 <= sem <= 53:
                sem = _semana_iso(ymd)
            evs.append((fi, ymd, estados.index(st), _num(celda(fila, "ph")), _num(celda(fila, "ps")), _num(celda(fila, "po")), _num(celda(fila, "larva")), sem))
            # Reporte diario: kilos (×10), color (índice; -1 sin color) y valores ×1e6 (-1 = celda vacía)
            col = _limpio(celda(fila, "color")) if ix["color"] >= 0 else ""
            if col and col not in colores:
                colores.append(col)
            post = [_num_o_nada(celda(fila, k)) for k in ("ph", "ps", "po")]
            post = None if all(p is None for p in post) else sum(p or 0 for p in post)   # = SUMA POSTURAS
            ind += [_redondeo(_num(celda(fila, "kilos")) * 10) if ix["kilos"] >= 0 else 0,
                    colores.index(col) if col else -1,
                    millon(_num_o_nada(celda(fila, "pic"))) if ix["pic"] >= 0 else -1,
                    millon(_num_o_nada(celda(fila, "prodi"))) if ix["prodi"] >= 0 else -1,
                    millon(_num_o_nada(celda(fila, "larva"))),
                    millon(post),
                    # 1 = entra al Reporte diario (PROCESO = FRESCO); si la base no trae PROCESO, entra todo
                    1 if ix["proceso"] < 0 or _clave_col(celda(fila, "proceso")) == PROCESO_REPORTE else 0]
        if not campos:
            error = "La base de fosfina no tiene filas con fecha y proveedor válidos."
            continue
        base, tope = min(e[1] for e in evs), max(e[1] for e in evs)
        plano = []
        for fi, ymd, st, ph, ps, po, larva, sem in evs:
            plano += [fi, _dias_entre(base, ymd), st, _redondeo(ph * 1e6), _redondeo(ps * 1e6), _redondeo(po * 1e6), _redondeo(larva * 1e6), sem]
        faltan_rep = [COLS_FOSFINA[k][1] for k in COLS_REPORTE if ix[k] < 0]
        if faltan_rep:
            log("Aviso (Reporte diario): la base de fosfina no trae " + ", ".join(faltan_rep) + "; esas tarjetas saldrán vacías.")
        # 'ind' va aparte de 'ev' para que fitosanitario.html siga leyendo igual que antes
        return {"nRows": n_filas, "maxYmd": tope, "base": base, "hasStatus": ix["status"] >= 0, "statuses": estados, "fields": campos, "ev": plano,
                "colores": colores, "ind": ind, "indW": IND_W, "indCols": [k for k in COLS_REPORTE if ix[k] >= 0],
                "headers": [" ".join(_txt(c).split()) for c in filas[r0] if _txt(c).strip()]}
    raise Detener(error)


COLS_RECEPCION = {
    "hora": (["HORACOSECHA"], "Hora Cosecha", True), "proceso": (["PROCESO"], "Proceso", True),
    "ua": (["UNIDADAGRICOLA"], "Unidad Agricola", True), "prov": (["PROVEEDOR"], "Proveedor", True),
    "campo": (["CAMPO"], "Campo", True), "clp": (["CLP"], "CLP", False),
    "kg": (["PESOPLANTA"], "Peso Planta", False), "cat": (["CATEGORIA"], "Categoria", False),
}


def leer_recepcion(ruta):
    """Recepción de materia prima, solo Proceso = FR. Devuelve {AAAAMMDD: {(ua, prov, campo, clp, cat): kilos}}."""
    error = "No encontré los encabezados Hora Cosecha y Proceso en el Excel de recepción."
    for _, filas in _hojas(ruta):
        enc = _encabezado(filas, COLS_RECEPCION, ["HORACOSECHA", "PROCESO"])
        if not enc:
            continue
        r0, ix, faltan = enc
        if faltan:
            error = "Faltan columnas en el Excel de recepción: " + ", ".join(faltan) + "."
            continue

        def celda(fila, k):
            return fila[ix[k]] if 0 <= ix[k] < len(fila) else None

        dias = {}
        for fila in filas[r0 + 1:]:
            ua, campo = _limpio(celda(fila, "ua")), _limpio(celda(fila, "campo"))
            if not ua and not campo:
                continue
            if _clave_col(celda(fila, "proceso")) != "FR":
                continue
            ymd = _ymd(celda(fila, "hora"))
            if not ymd:
                continue
            llave = (ua, _limpio(celda(fila, "prov")), campo, _limpio(celda(fila, "clp")), _limpio(celda(fila, "cat")))
            dia = dias.setdefault(ymd, {})
            dia[llave] = dia.get(llave, 0.0) + _num(celda(fila, "kg"))
        return dias
    raise Detener(error)


# ============================================================
# 1b. CALIDADES Y CALIBRES + LONGITUDES (Reporte diario)
#     Se guardan solo TOTALES por día y unidad agrícola (no cada fila), con el mismo filtro del
#     Power BI: PROCESO = FRESCO y STATUS = EVALUADO. Así el archivo queda chico aunque la base pese 50 MB.
# ============================================================

NOMBRE_CALIDADES = os.environ.get("CALIDADES_NOMBRE", "CALIDADES Y CALIBRES")   # parte del nombre en Drive
NOMBRE_LONGITUDES = os.environ.get("LONGITUDES_NOMBRE", "LONGITUDES")
STATUS_REPORTE = "EVALUADO"

# Mismas listas que las medidas DAX de Johan (se comparan sin espacios ni mayúsculas)
CAL_C_RAM = {"C-", "RAMEADO"}                 # Kg C- y Rameados: CALIDAD IN {"C -", "Rameado"}
CAL_SMALL = {"4-5.9MM", "6-7.9"}              # Kg Small: CALIBRE IN {"4 - 5.9 mm", "6 - 7.9"}
# *LONG (año 2026): el texto se compara EXACTO, porque ">28 cm" (sin espacio) no entra y "> 28 cm" sí
LONG_RANGO = {"<10 cm": 0, "<14 cm": 0, "10 - 13.9 cm": 0, "14 - 17.9 cm": 0, "18 - 24 cm": 0, "24.1 - 26 cm": 0,
              "26.1 - 28 cm": 1, "28.1 - 30 cm": 2, "> 28 cm": 2, "> 30 cm": 3}

COLS_CAL = {
    "fecha": (["FECHACOSECHA", "FECHA"], "FECHA COSECHA", True), "sem": (["SEMANA"], "SEMANA", False),
    "ua": (["UNIDADAGRICOLA"], "UNIDAD AGRICOLA", True), "proceso": (["PROCESO"], "PROCESO", True),
    "status": (["STATUS"], "STATUS", True), "calidad": (["CALIDAD"], "CALIDAD", True),
    "calibre": (["CALIBRE"], "CALIBRE", True), "kg": (["VALORPLANTAKG"], "Valor Planta Kg", True),
    "car": (["CARACTERISTICA"], "CARACTERISTICA", False),
}
COLS_LON = {
    "fecha": (["FECHACOSECHA", "FECHA"], "FECHA COSECHA", True), "sem": (["SEMANA"], "SEMANA", False),
    "ua": (["UNIDADAGRICOLA"], "UNIDAD AGRICOLA", True), "proceso": (["PROCESO"], "PROCESO", True),
    "status": (["STATUS"], "STATUS", True), "lon": (["LONGITUD"], "LONGITUD", True),
    "pct": (["PORCENTAJE"], "PORCENTAJE", True),
    "car": (["CARACTERISTICA"], "CARACTERISTICA", False),
}


def _filas_stream(ruta, preferidas=()):
    """(hoja, iterador de filas). Con python-calamine (mucho más rápido en bases grandes) o, si no está, openpyxl."""
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        CalamineWorkbook = None
    if CalamineWorkbook:
        wb = CalamineWorkbook.from_path(str(ruta))
        for n in sorted(wb.sheet_names, key=lambda n: 0 if _clave_col(n) in preferidas else 1):
            hoja = wb.get_sheet_by_name(n)
            yield n, (hoja.iter_rows() if hasattr(hoja, "iter_rows") else iter(hoja.to_python()))
        return
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)
    for n in sorted(wb.sheetnames, key=lambda n: 0 if _clave_col(n) in preferidas else 1):
        yield n, wb[n].iter_rows(values_only=True)


def _tabla(ruta, cols, obligatorias, preferidas, que):
    """Busca el encabezado en las primeras filas y devuelve (índices de columna, iterador del resto de filas)."""
    for _, it in _filas_stream(ruta, preferidas):
        cabeza = []
        for fila in it:
            cabeza.append(list(fila))
            if len(cabeza) >= 20:
                break
        enc = _encabezado(cabeza, cols, obligatorias)
        if not enc:
            continue
        r0, ix, faltan = enc
        if faltan:
            raise ValueError(f"faltan columnas en la base de {que}: " + ", ".join(faltan))

        def filas(cabeza=cabeza, r0=r0, it=it):
            yield from cabeza[r0 + 1:]
            yield from it
        return ix, filas()
    raise ValueError(f"no encontré los encabezados de la base de {que}")


def _semana_de(v, ymd):
    m = re.search(r"\d+", _txt(v))
    s = int(m.group(0)) if m else 0
    return s if 1 <= s <= 53 else _semana_iso(ymd)


def _pasa_filtro(f, ix, caracteristica):
    """Filtros internos del Power BI (no se muestran en la página):
    PROCESO = FRESCO, STATUS = EVALUADO y, según la base, CARACTERISTICA = NORMAL
    (calidades: C- y Ram / Small) o CARACTERISTICA no vacía (longitudes)."""
    def c(k):
        return f[ix[k]] if 0 <= ix[k] < len(f) else None
    if _clave_col(c("proceso")) != PROCESO_REPORTE or _clave_col(c("status")) != STATUS_REPORTE:
        return False
    car = _clave_col(c("car"))
    if caracteristica == "NORMAL":
        return ix["car"] < 0 or car == "NORMAL"
    return ix["car"] < 0 or car != ""


def _empacar_dia_ua(tot, ancho, escala):
    """{(ymd, sem, ua): [valores]} -> {'base', 'uas', 'rows'} con filas planas [día, semana, ua, valores×escala...]."""
    if not tot:
        return None
    base = min(k[0] for k in tot)
    uas, pos, filas = [], {}, []
    for (ymd, sem, ua), vals in sorted(tot.items()):
        if ua not in pos:
            pos[ua] = len(uas)
            uas.append(ua)
        filas += [_dias_entre(base, ymd), sem, pos[ua]] + [_redondeo(v * escala) for v in vals]
    return {"base": base, "maxYmd": max(k[0] for k in tot), "uas": uas, "w": 3 + ancho, "rows": filas}


def leer_calidades(ruta):
    """Por día y unidad agrícola: kilos totales, kilos C- y Rameados y kilos Small (Valor Planta Kg)."""
    ix, filas = _tabla(ruta, COLS_CAL, ["CALIDAD", "CALIBRE"], ("CALIDADESYCALIBRES",), "calidades y calibres")
    tot, n = {}, 0
    for f in filas:
        if not f or not _pasa_filtro(f, ix, "NORMAL"):
            continue
        ymd = _ymd(f[ix["fecha"]])
        if not ymd:
            continue
        kg = _num(f[ix["kg"]])
        if not kg:
            continue
        n += 1
        llave = (ymd, _semana_de(f[ix["sem"]] if ix["sem"] >= 0 else None, ymd), _limpio(f[ix["ua"]]))
        t = tot.setdefault(llave, [0.0, 0.0, 0.0])
        t[0] += kg
        if _clave_col(f[ix["calidad"]]) in CAL_C_RAM:
            t[1] += kg
        if _clave_col(f[ix["calibre"]]) in CAL_SMALL:
            t[2] += kg
    log(f"Calidades y calibres: {n} filas FRESCO/EVALUADO en {len({k[0] for k in tot})} días.")
    return _empacar_dia_ua(tot, 3, 100)        # kilos ×100


def leer_longitudes(ruta):
    """Por día y unidad agrícola: suma de PORCENTAJE de cada rango (<26, 26.1-28, 28.1-30, >30 cm)."""
    ix, filas = _tabla(ruta, COLS_LON, ["LONGITUD", "PORCENTAJE"], ("BD", "LONGITUDES"), "longitudes")
    tot, n = {}, 0
    for f in filas:
        if not f or not _pasa_filtro(f, ix, "NO_VACIA"):
            continue
        ymd = _ymd(f[ix["fecha"]])
        rango = LONG_RANGO.get(" ".join(_txt(f[ix["lon"]]).split()))
        if not ymd or rango is None or ymd // 10000 < 2026:      # años anteriores usan otros rangos en el DAX
            continue
        n += 1
        llave = (ymd, _semana_de(f[ix["sem"]] if ix["sem"] >= 0 else None, ymd), _limpio(f[ix["ua"]]))
        tot.setdefault(llave, [0.0, 0.0, 0.0, 0.0])[rango] += _num(f[ix["pct"]])
    log(f"Longitudes: {n} filas FRESCO/EVALUADO en {len({k[0] for k in tot})} días.")
    return _empacar_dia_ua(tot, 4, 1e4)        # suma de porcentajes ×1e4


# ============================================================
# 2. EMPAQUETADO (el mismo formato que la página ya sabe leer)
# ============================================================

def empacar_recepcion(dias):
    """{AAAAMMDD: {(ua, prov, campo, clp, cat): kg}} -> bloque 'rec' compacto."""
    if not dias:
        return None
    base, tope = min(dias), max(dias)
    listas = {k: [] for k in ("ua", "prov", "campo", "clp", "cat")}
    pos = {k: {} for k in listas}

    def indice(k, v):
        if v not in pos[k]:
            pos[k][v] = len(listas[k])
            listas[k].append(v)
        return pos[k][v]

    filas = []
    for ymd in sorted(dias):
        for (ua, prov, campo, clp, cat), kg in sorted(dias[ymd].items()):
            filas += [_dias_entre(base, ymd), indice("ua", ua), indice("prov", prov), indice("campo", campo), indice("clp", clp), _redondeo(kg * 10), indice("cat", cat)]
    n = len(filas) // 7
    return {"nRows": n, "nFR": n, "maxYmd": tope, "base": base, **listas, "rows": filas}


def desempacar_recepcion(rec):
    dias = {}
    if not rec:
        return dias
    f = rec["rows"]
    for i in range(0, len(f), 7):
        ymd = _sumar_dias(rec["base"], f[i])
        llave = (rec["ua"][f[i + 1]], rec["prov"][f[i + 2]], rec["campo"][f[i + 3]], rec["clp"][f[i + 4]], rec["cat"][f[i + 6]])
        dias.setdefault(ymd, {})[llave] = f[i + 5] / 10
    return dias


def combinar_recepcion(guardados, nuevos, objetivo):
    """En cada día de 'objetivo' se reemplazan las unidades agrícolas que trae el Excel (igual que el Flujo de cosecha)."""
    total = 0
    for ymd in sorted(nuevos):
        if objetivo is not None and ymd not in objetivo:
            continue
        uas = {k[0] for k in nuevos[ymd]}
        dia = {k: v for k, v in guardados.get(ymd, {}).items() if k[0] not in uas}
        dia.update(nuevos[ymd])
        guardados[ymd] = dia
        total += len(nuevos[ymd])
    return total


# ============================================================
# 3. CIFRADO (mismo esquema que datos.json: una cerradura por DNI + la del administrador)
# ============================================================

def _b64(b):
    return base64.b64encode(b).decode("ascii")


def _clave(texto, sal, iteraciones):
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
    return PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=sal, iterations=iteraciones).derive(texto.encode("utf-8"))


def _llave_dni(v):
    d = re.sub(r"[^0-9A-Z]", "", str(v or "").upper())
    return re.sub(r"^0+(?=\d)", "", d) if d.isdigit() else d


def _abrir_adm(pk, clave_admin):
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    adm = pk["adm"]
    k = _clave("adm:" + clave_admin, base64.b64decode(adm["s"]), int(pk["it"]))
    return json.loads(AESGCM(k).decrypt(base64.b64decode(adm["i"]), base64.b64decode(adm["c"]), None))


def accesos(clave_admin):
    """Lee de datos.json quiénes tienen acceso: (lista de DNIs, iteraciones, sal). No toca datos.json."""
    from cryptography.exceptions import InvalidTag
    if not RUTA_DATOS.exists():
        raise Detener("No existe datos.json: primero tiene que estar publicado el Flujo de cosecha.")
    pk = json.loads(RUTA_DATOS.read_text(encoding="utf-8"))
    if not pk.get("adm"):
        raise Detener("datos.json no tiene sección de administrador.")
    try:
        a = _abrir_adm(pk, clave_admin)
    except InvalidTag:
        raise Detener("La clave de administrador (secreto ADMIN_CLAVE) no coincide con la del Flujo de cosecha.")
    return list(a.get("dnis") or []), int(pk["it"]), base64.b64decode(pk["salt"])


def abrir_fito(clave_admin):
    """Devuelve (contenido, firma) del fito.json que ya existe, o (None, None)."""
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    if not RUTA_FITO.exists():
        return None, None
    try:
        pk = json.loads(RUTA_FITO.read_text(encoding="utf-8"))
        a = _abrir_adm(pk, clave_admin)
        plano = AESGCM(base64.b64decode(a["k"])).decrypt(base64.b64decode(pk["iv"]), base64.b64decode(pk["data"]), None)
        return json.loads(gzip.decompress(plano) if pk.get("z") else plano), a.get("firma")
    except Exception as e:   # archivo dañado o cifrado con otra clave: se rehace desde cero
        log(f"Aviso: no se pudo abrir el fito.json anterior ({type(e).__name__}); se arma uno nuevo.")
        return None, None


def _firma(dnis, sal, iteraciones):
    """Cambia si cambia la lista de accesos: obliga a volver a cifrar aunque los datos sean los mismos."""
    return hashlib.sha256(json.dumps([sorted(_llave_dni(d) for d in dnis), _b64(sal), iteraciones]).encode("utf-8")).hexdigest()


def guardar_fito(contenido, dnis, clave_admin, iteraciones, sal):
    """Las cerraduras usan la MISMA sal de datos.json: así la sesión iniciada en el Flujo de cosecha también abre este archivo."""
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    plano = gzip.compress(json.dumps(contenido, ensure_ascii=False, separators=(",", ":")).encode("utf-8"), mtime=0)
    cruda, iv = os.urandom(32), os.urandom(12)
    datos = AESGCM(cruda).encrypt(iv, plano, None)
    llaves = []
    for dni in dnis:
        kiv = os.urandom(12)
        llaves.append({"i": _b64(kiv), "c": _b64(AESGCM(_clave(_llave_dni(dni), sal, iteraciones)).encrypt(kiv, cruda, None))})
    s_adm, i_adm = os.urandom(16), os.urandom(12)
    c_adm = AESGCM(_clave("adm:" + clave_admin, s_adm, iteraciones)).encrypt(
        i_adm, json.dumps({"k": _b64(cruda), "firma": _firma(dnis, sal, iteraciones)}, separators=(",", ":")).encode("utf-8"), None)
    pk = {"v": 1, "gen": contenido.get("gen"), "salt": _b64(sal), "it": iteraciones, "keys": llaves, "iv": _b64(iv), "z": True,
          "data": _b64(datos), "adm": {"s": _b64(s_adm), "i": _b64(i_adm), "c": _b64(c_adm)}}
    RUTA_FITO.write_text(json.dumps(pk, separators=(",", ":")), encoding="utf-8")


# ============================================================
# 4. DESCARGA DE LA BASE DE FOSFINA DESDE GOOGLE DRIVE (como bd.py)
# ============================================================

def descargar_fosfina(carpeta, buscar=None, archivo="fosfina.xlsx"):
    """Baja de la carpeta de Drive el archivo más reciente cuyo nombre contenga `buscar` (por defecto NOMBRE_FOSFINA).
    Devuelve (ruta, nombre) o (None, None)."""
    buscar = buscar or NOMBRE_FOSFINA
    cred_txt = os.environ.get("GOOGLE_CREDENCIALES", "").strip()
    carpeta_id = os.environ.get("DRIVE_CARPETA_ID", "").strip()
    if not cred_txt or not carpeta_id:
        log("Aviso: faltan los secretos GOOGLE_CREDENCIALES y/o DRIVE_CARPETA_ID; no se descarga «" + buscar + "».")
        return None, None
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload

    creds = Credentials.from_service_account_info(json.loads(cred_txt), scopes=["https://www.googleapis.com/auth/drive.readonly"])
    servicio = build("drive", "v3", credentials=creds, cache_discovery=False)
    nombre = buscar.replace("'", "\\'")
    res = servicio.files().list(
        q=f"'{carpeta_id}' in parents and name contains '{nombre}' and trashed = false",
        fields="files(id, name, mimeType, modifiedTime)", orderBy="modifiedTime desc", pageSize=5,
        supportsAllDrives=True, includeItemsFromAllDrives=True).execute()
    archivos = res.get("files", [])
    if not archivos:
        log(f"Aviso: en la carpeta de Drive no hay ningún archivo cuyo nombre contenga '{buscar}'.")
        return None, None
    info = archivos[0]
    log(f"Drive: '{info['name']}' (modificado {info.get('modifiedTime')}).")
    if info.get("mimeType") == "application/vnd.google-apps.spreadsheet":   # hoja de Google: se exporta como Excel
        pedido = servicio.files().export_media(fileId=info["id"], mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        pedido = servicio.files().get_media(fileId=info["id"], supportsAllDrives=True)
    buffer = io.BytesIO()
    bajada, listo = MediaIoBaseDownload(buffer, pedido), False
    while not listo:
        _, listo = bajada.next_chunk()
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / archivo
    ruta.write_bytes(buffer.getvalue())
    return ruta, info["name"]


# ============================================================
# 5. PROGRAMA
# ============================================================

def main():
    ap = argparse.ArgumentParser(description="Actualiza fito.json (recepción FR + base de fosfina) para fitosanitario.html")
    ap.add_argument("--dias", type=int, default=int(os.environ.get("DIAS_ATRAS") or 1), help="cuántos días hacia atrás trae el Excel de recepción (1 = solo ayer)")
    ap.add_argument("--recepcion", help="usar este Excel de recepción en vez del que bajó actualizar.py")
    ap.add_argument("--todo", action="store_true", help="con --recepcion: guardar TODOS los días que traiga el Excel (para cargar historia)")
    ap.add_argument("--fosfina", help="usar este Excel de fosfina en vez de bajarlo de Drive")
    ap.add_argument("--calidades", help="usar este Excel de calidades y calibres en vez de bajarlo de Drive")
    ap.add_argument("--longitudes", help="usar este Excel de longitudes en vez de bajarlo de Drive")
    ap.add_argument("--hoy", help="fecha de hoy AAAA-MM-DD (solo para pruebas)")
    ap.add_argument("--incluir-hoy", action="store_true", help="guardar también lo que va de hoy (día incompleto)")
    a = ap.parse_args()

    clave_admin = os.environ.get("ADMIN_CLAVE", "")
    if not clave_admin:
        raise Detener("Falta el secreto ADMIN_CLAVE.")
    dnis, it, sal = accesos(clave_admin)
    anterior, firma_anterior = abrir_fito(clave_admin)
    anterior = anterior or {}
    carpeta = Path(os.environ.get("DESCARGAS_DIR", "descargas")).resolve()

    # --- recepción: los mismos días que publica el Flujo de cosecha (desde Fecha Inicio hasta ayer; en la tarde, también hoy)
    hoy = datetime.strptime(a.hoy, "%Y-%m-%d").date() if a.hoy else datetime.now(LIMA).date()
    dias = max(1, min(a.dias, 31))
    incluir_hoy = a.incluir_hoy or os.environ.get("INCLUIR_HOY", "").strip().lower() in ("1", "true", "si", "sí")
    objetivo = None if (a.recepcion and a.todo) else {int((hoy - timedelta(days=dias) + timedelta(days=i)).strftime("%Y%m%d")) for i in range(dias + (1 if incluir_hoy else 0))}
    ruta_rec = Path(a.recepcion) if a.recepcion else None
    if not ruta_rec:
        excels = [p for p in carpeta.glob("*.xls*") if p.name != "fosfina.xlsx"] if carpeta.exists() else []
        ruta_rec = max(excels, key=lambda p: p.stat().st_mtime) if excels else None
    rec_dias = desempacar_recepcion(anterior.get("rec"))
    if ruta_rec:
        nuevos = leer_recepcion(ruta_rec)
        n = combinar_recepcion(rec_dias, nuevos, objetivo)
        usados = sorted(d for d in nuevos if objetivo is None or d in objetivo)
        log(f"Recepción (FR): {n} campos-día en {len(usados)} día(s)" + (f", del {str(usados[0])[6:]}/{str(usados[0])[4:6]} al {str(usados[-1])[6:]}/{str(usados[-1])[4:6]}." if usados else "."))
    else:
        log("Aviso: no hay Excel de recepción de esta corrida; se conserva la recepción ya guardada.")
    rec = empacar_recepcion(rec_dias)

    # --- fosfina: se reemplaza completa
    fito, nombre_fito = None, None
    try:
        ruta_fos, nombre_fito = (Path(a.fosfina), Path(a.fosfina).stem) if a.fosfina else descargar_fosfina(carpeta)
        if ruta_fos:
            fito = leer_fosfina(ruta_fos)
            log(f"Fosfina: {fito['nRows']} evaluaciones de {len(fito['fields'])} campos, hasta el {str(fito['maxYmd'])[6:]}/{str(fito['maxYmd'])[4:6]}/{str(fito['maxYmd'])[:4]}.")
    except Detener:
        raise
    except Exception as e:
        log(f"Aviso: no se pudo actualizar la base de fosfina ({type(e).__name__}: {e}). Se conserva la anterior.")
    if fito is None:
        fito = anterior.get("fito")
    if not fito or not rec:
        raise Detener("Todavía falta " + ("la base de fosfina" if not fito else "la recepción") + ": no se publica fito.json.")
    fito["name"] = "Base fitosanitaria (Drive)"
    rec["name"] = "Recepción de materia prima (intranet)"

    # --- calidades y longitudes (Reporte diario): si algo falla, se conserva lo anterior y lo demás sigue
    extra = {}
    for clave, buscar, archivo, leer, local in (("cal", NOMBRE_CALIDADES, "calidades.xlsx", leer_calidades, a.calidades),
                                                 ("lon", NOMBRE_LONGITUDES, "longitudes.xlsx", leer_longitudes, a.longitudes)):
        try:
            ruta = Path(local) if local else descargar_fosfina(carpeta / "drive", buscar, archivo)[0]
            extra[clave] = leer(ruta) if ruta else None
        except Exception as e:
            log(f"Aviso: no se pudo leer la base de {buscar} ({type(e).__name__}: {e}). Se conserva la anterior.")
            extra[clave] = None
        if extra[clave] is None:
            extra[clave] = anterior.get(clave)
        if extra[clave] is None:
            del extra[clave]

    contenido = {"v": 1, "fito": fito, "rec": rec, **extra}
    # bloques que se suben desde la laptop (aquí solo se conservan):
    #   rsql = kilos recepcionados del SQL (recepcion_sql.py) · dren = kilos drenados del Power BI (drenados_pbi.py)
    for k in ("rsql", "dren"):
        if anterior.get(k):
            contenido[k] = anterior[k]
    igual = {k: anterior.get(k) for k in ("v", "fito", "rec", "cal", "lon", "rsql", "dren") if k in anterior or k in contenido} == contenido
    if igual and firma_anterior == _firma(dnis, sal, it):
        log("fito.json ya estaba al día: no se cambia nada.")
        return 0
    contenido["gen"] = datetime.now(LIMA).strftime("%Y-%m-%dT%H:%M")
    guardar_fito(contenido, dnis, clave_admin, it, sal)
    log(f"fito.json actualizado: {rec['nRows']} campos-día de recepción en {len(rec_dias)} días; {fito['nRows']} evaluaciones; {len(dnis)} DNIs con acceso.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Detener as e:
        log(f"DETENIDO: {e}")
        sys.exit(1)
