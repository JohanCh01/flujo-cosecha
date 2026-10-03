"""
Actualización automática del Flujo Comparativo de la Cosecha.

Qué hace, en orden:
1. Entra a la intranet de Danper, abre el reporte "Recepción de Materia Prima",
   elige ESPÁRRAGO / ESPÁRRAGO VERDE, pone Fecha Inicio = ayer y Fecha Fin = hoy,
   espera a que cargue y exporta el Excel.
2. Lee el Excel igual que la página web (mismas columnas y mismas reglas).
3. Comprueba que las fechas del Excel correspondan a lo que se pidió. Si no
   coinciden (día y mes invertidos), repite la descarga escribiendo la fecha
   en el otro orden; si aun así no coinciden, se detiene SIN publicar.
4. Abre datos.json (los datos cifrados del enlace), reemplaza los días
   descargados, vuelve a cifrar con la misma lista de DNIs y guarda.

Pensado para correr en GitHub Actions (ver .github/workflows/actualizar.yml),
pero también corre en una PC con Python:

    python actualizar.py --solo-descarga --visible     (prueba: solo descarga y resume)
    python actualizar.py --excel reporte.xlsx          (usa un Excel ya descargado)

Credenciales: nunca van escritas aquí. Se leen de variables de entorno
(INTRANET_USUARIO, INTRANET_CLAVE, ADMIN_CLAVE); en GitHub son "secrets".
"""

import argparse
import base64
import gzip
import json
import math
import os
import re
import sys
import time
import unicodedata
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

LIMA = timezone(timedelta(hours=-5))
URL_LOGIN = "https://intranet.danper.com/#/login"
URL_REPORTE = "https://intranet.danper.com/#/recojoMP/reports/reporteRecepcionMateriaPrima"
TIPO_CULTIVO = "ESPÁRRAGO"
MATERIA_PRIMA = "ESPÁRRAGO VERDE"
RUTA_DATOS = Path(os.environ.get("DATOS_JSON", "datos.json"))
EPOCA = datetime(1970, 1, 1)


def log(msg):
    print(f"[{datetime.now(LIMA).strftime('%H:%M:%S')}] {msg}", flush=True)


class Detener(Exception):
    """Error esperado: se muestra el mensaje y no se publica nada."""


class SinOpciones(Detener):
    """Una lista desplegable no mostró opciones (todavía no cargan o no se desplegó)."""


# ============================================================
# 1. LECTURA DEL EXCEL (mismas reglas que la página web)
# ============================================================

COLS = ["ua", "prov", "campo", "cant", "c", "r", "s", "p", "placa", "guia", "ticket", "alm", "pr", "kp", "kc", "kb"]
CAMPOS = {
    "r": ["fecha recojo", "fecha de recojo", "hora recojo"],
    "c": ["hora cosecha", "inicio de cosecha", "inicio cosecha"],
    "s": ["hora salida", "salida"],
    "p": ["hora recepcion", "recepcion"],
    "ua": ["unidad agricola"], "prov": ["proveedor"], "campo": ["campo"], "cant": ["cantidad", "cant"],
    "placa": ["placa"], "guia": ["guia remision", "guia de remision", "guia"],
    "ticket": ["nro ticket", "ticket", "numero ticket"], "alm": ["almacen"], "pr": ["proceso"],
    "kp": ["peso planta"], "kc": ["peso campo"], "kb": ["peso bruto"],
}
OBLIGATORIAS = {"r": "Fecha Recojo", "ua": "Unidad Agricola", "prov": "Proveedor", "campo": "Campo", "cant": "Cantidad"}
RE_FECHA = re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})(?:[ T]+(\d{1,2}):(\d{2})(?::\d{2})?\s*(?:([ap])\.?\s*m\.?)?)?$", re.I)
RE_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:[ T]+(\d{1,2}):(\d{2}))?")


def _norm(s):
    s = "" if s is None else str(s)
    s = "".join(ch for ch in unicodedata.normalize("NFD", s) if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _minutos(y, mo, d, h=0, mi=0):
    return int((datetime(y, mo, d, h, mi) - EPOCA).total_seconds() // 60)


def _fecha_hora(v):
    """Devuelve minutos desde 1970 (hora local del reporte, sin zona) o None."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return _minutos(v.year, v.month, v.day, v.hour, v.minute)
    if isinstance(v, date):
        return _minutos(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(math.floor((v - 25569) * 1440 + 0.5)) if 20000 < v < 90000 else None
    s = str(v).strip()
    m = RE_FECHA.match(s)
    h = mi = 0
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        if m.group(4) is not None:
            h, mi = int(m.group(4)), int(m.group(5))
            if m.group(6):
                pm = m.group(6).lower() == "p"
                if pm and h < 12:
                    h += 12
                if not pm and h == 12:
                    h = 0
    else:
        m = RE_ISO.match(s)
        if not m:
            return None
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if m.group(4) is not None:
            h, mi = int(m.group(4)), int(m.group(5))
    if not (1 <= mo <= 12 and 1 <= d <= 31 and h <= 23 and mi <= 59):
        return None
    try:
        return _minutos(y, mo, d, h, mi)
    except ValueError:
        return None


def _txt(v):
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def _entero_si_puede(n):
    return int(n) if float(n).is_integer() else n


def _num(v):
    if v is None or v == "":
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(n):
        return None
    return _entero_si_puede(math.floor(n * 100 + 0.5) / 100)


def _cant(v):
    try:
        n = float(v)
    except (TypeError, ValueError):
        return 0
    return _entero_si_puede(n) if math.isfinite(n) else 0


def dia_de(minutos):
    return (EPOCA + timedelta(minutes=minutos)).strftime("%Y-%m-%d")


def leer_excel(ruta):
    """Devuelve la lista de recojos del Excel (cada uno un dict con las claves de COLS)."""
    from openpyxl import load_workbook

    wb = load_workbook(ruta, read_only=True, data_only=True)
    faltan = None
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        for hi in range(min(len(filas), 25)):
            cab = [_norm(x) for x in (filas[hi] or [])]
            idx = {}
            for k, alias in CAMPOS.items():
                for a in alias:
                    if a in cab:
                        idx[k] = cab.index(a)
                        break
            if "r" not in idx:
                continue
            f = [OBLIGATORIAS[k] for k in OBLIGATORIAS if k not in idx]
            if f:
                faltan = f
                continue
            salida = []
            for w in filas[hi + 1:]:
                w = w or []

                def g(k):
                    return w[idx[k]] if k in idx and idx[k] < len(w) else None

                r, ua = _fecha_hora(g("r")), _txt(g("ua"))
                if r is None or not ua:
                    continue   # fila de totales o vacía
                salida.append({
                    "ua": ua, "prov": _txt(g("prov")) or "Sin proveedor", "campo": _txt(g("campo")) or "Sin campo",
                    "cant": _cant(g("cant")), "c": _fecha_hora(g("c")), "r": r, "s": _fecha_hora(g("s")), "p": _fecha_hora(g("p")),
                    "placa": _txt(g("placa")), "guia": _txt(g("guia")), "ticket": _txt(g("ticket")), "alm": _txt(g("alm")),
                    "pr": None if "pr" not in idx else (_txt(g("pr")) or "(sin proceso)"),
                    "kp": _num(g("kp")), "kc": _num(g("kc")), "kb": _num(g("kb")),
                })
            return salida
    raise Detener("Al Excel le faltan las columnas: " + ", ".join(faltan) if faltan else "El Excel no trae la columna Fecha Recojo.")


def empacar(dia, filas):
    doc = {"fecha": dia, "n": len(filas)}
    for c in COLS:
        doc[c] = [f.get(c) for f in filas]
    return doc


def desempacar(doc):
    n = len(doc.get("r") or [])
    filas = []
    for i in range(n):
        f = {c: (doc[c][i] if doc.get(c) else None) for c in COLS}
        if isinstance(f["r"], (int, float)):
            filas.append(f)
    return filas


# ============================================================
# 2. datos.json: abrir, combinar y volver a cifrar
#    (mismo formato que usa la página: PBKDF2-SHA256 + AES-GCM)
# ============================================================

def _b64(b):
    return base64.b64encode(b).decode("ascii")


def _clave(texto, sal, iteraciones):
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
    return PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=sal, iterations=iteraciones).derive(texto.encode("utf-8"))


def _llave_dni(v):
    d = re.sub(r"[^0-9A-Z]", "", str(v or "").upper())
    return re.sub(r"^0+(?=\d)", "", d) if d.isdigit() else d


def abrir_datos(clave_admin):
    """Descifra datos.json con la clave de administrador. Devuelve (contenido, lista_dnis, iteraciones)."""
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    from cryptography.exceptions import InvalidTag

    if not RUTA_DATOS.exists():
        raise Detener("No existe datos.json en el repositorio. Primero entra al enlace como administrador y carga al menos un Excel.")
    pk = json.loads(RUTA_DATOS.read_text(encoding="utf-8"))
    if not pk.get("adm"):
        raise Detener("datos.json no tiene sección de administrador. Vuelve a publicar desde el enlace (Accesos → Guardar y publicar).")
    it = int(pk["it"])
    adm = pk["adm"]
    try:
        k_adm = _clave("adm:" + clave_admin, base64.b64decode(adm["s"]), it)
        a = json.loads(AESGCM(k_adm).decrypt(base64.b64decode(adm["i"]), base64.b64decode(adm["c"]), None))
    except InvalidTag:
        raise Detener("La clave de administrador (secreto ADMIN_CLAVE) no coincide con la que usas en el enlace.")
    plano = AESGCM(base64.b64decode(a["k"])).decrypt(base64.b64decode(pk["iv"]), base64.b64decode(pk["data"]), None)
    if pk.get("z"):
        plano = gzip.decompress(plano)
    return json.loads(plano), list(a.get("dnis") or []), it


def guardar_datos(contenido, dnis, clave_admin, iteraciones, hoy):
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    plano = gzip.compress(json.dumps(contenido, ensure_ascii=False, separators=(",", ":")).encode("utf-8"), mtime=0)
    cruda = os.urandom(32)
    iv = os.urandom(12)
    datos = AESGCM(cruda).encrypt(iv, plano, None)
    sal = os.urandom(16)
    llaves = []
    for dni in dnis:   # una "cerradura" por cada DNI autorizado
        kiv = os.urandom(12)
        llaves.append({"i": _b64(kiv), "c": _b64(AESGCM(_clave(_llave_dni(dni), sal, iteraciones)).encrypt(kiv, cruda, None))})
    s_adm, i_adm = os.urandom(16), os.urandom(12)
    c_adm = AESGCM(_clave("adm:" + clave_admin, s_adm, iteraciones)).encrypt(
        i_adm, json.dumps({"k": _b64(cruda), "dnis": dnis}, separators=(",", ":")).encode("utf-8"), None)
    dias = contenido["index"]["dias"]
    pk = {"v": 1, "gen": hoy.strftime("%Y-%m-%d"), "desde": dias[0] if dias else None, "hasta": dias[-1] if dias else None,
          "salt": _b64(sal), "it": iteraciones, "keys": llaves, "iv": _b64(iv), "z": True, "data": _b64(datos),
          "adm": {"s": _b64(s_adm), "i": _b64(i_adm), "c": _b64(c_adm)}}
    RUTA_DATOS.write_text(json.dumps(pk, separators=(",", ":")), encoding="utf-8")


def combinar(contenido, por_dia, etiqueta):
    """Reemplaza, en cada día descargado, las unidades agrícolas que trae el Excel (igual que Cargar Excel)."""
    dias = contenido.setdefault("days", {})
    indice = contenido.setdefault("index", {"dias": [], "uas": [], "cargas": []})
    total = 0
    for dia, nuevas in sorted(por_dia.items()):
        previas = desempacar(dias[dia]) if dia in dias else []
        uas = {f["ua"] for f in nuevas}
        dias[dia] = empacar(dia, [f for f in previas if f["ua"] not in uas] + nuevas)
        total += len(nuevas)
    claves = sorted(por_dia)
    indice["dias"] = sorted(set(indice.get("dias", [])) | set(claves))
    indice["uas"] = sorted(set(indice.get("uas", [])) | {f["ua"] for fs in por_dia.values() for f in fs})
    nueva = {"nombre": etiqueta, "filas": total, "desde": claves[0], "hasta": claves[-1],
             "cuando": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")}
    cargas = [c for c in indice.get("cargas", []) if not (c.get("nombre") == nueva["nombre"] and c.get("desde") == nueva["desde"] and c.get("hasta") == nueva["hasta"])]
    indice["cargas"] = (cargas + [nueva])[-40:]
    return total


# ============================================================
# 3. DESCARGA DESDE LA INTRANET (Selenium)
# ============================================================

MENU = ["Gestión Agrícola", "Recojo Materia Prima", "Reportes", "Recepción Materia Prima"]

# Busca un control (lista o casilla) por el texto de su recuadro, sin depender de cómo esté armada la página.
JS_CONTROL = """
var etiqueta = arguments[0], tipo = arguments[1];
function plano(s){ return (s||'').normalize('NFD').replace(/[\\u0300-\\u036f]/g,'').toLowerCase().replace(/\\s+/g,' ').trim(); }
function visible(e){ return !!(e && e.getClientRects().length); }
var q = plano(etiqueta), mejor = null;
var lista = Array.prototype.slice.call(document.querySelectorAll(tipo === 'lista' ? 'mat-select, select, [role="combobox"]' : 'input'));
for (var i = 0; i < lista.length; i++) {
  var e = lista[i];
  if (!visible(e)) continue;
  var caja = e.closest('mat-form-field') || e.parentElement;
  var propios = plano([e.getAttribute('aria-label'), e.getAttribute('placeholder'), e.getAttribute('data-placeholder'), e.getAttribute('formcontrolname'), e.getAttribute('name')].join(' '));
  var texto = plano(caja ? caja.innerText : '');
  if (texto.indexOf(q) === 0) return e;                      // el recuadro empieza con la etiqueta
  if (!mejor && (texto.indexOf(q) >= 0 || propios.indexOf(q) >= 0 || propios.replace(/ /g,'').indexOf(q.replace(/ /g,'')) >= 0)) mejor = e;
}
return mejor;
"""

JS_OPCIONES = """
var r = [], ops = document.querySelectorAll('mat-option, [role="option"], .cdk-overlay-container li');
for (var i = 0; i < ops.length; i++) {
  var o = ops[i];
  if (!o.getClientRects().length) continue;
  var t = (o.innerText || o.textContent || '').replace(/\\s+/g,' ').trim();
  if (t) r.push([o, t]);
}
return r;
"""

JS_PISTAS = """
function corto(s){ return (s||'').replace(/\\s+/g,' ').trim().slice(0,70); }
function visible(e){ return !!e.getClientRects().length; }
var r = {listas:[], casillas:[], opciones:[], capas:0, paginador:''};
Array.prototype.slice.call(document.querySelectorAll('mat-select, select, [role="combobox"]')).filter(visible).slice(0,10).forEach(function(e){
  var c = e.closest('mat-form-field') || e.parentElement;
  r.listas.push(e.tagName.toLowerCase() + (e.getAttribute('aria-disabled') === 'true' ? ' (desactivada)' : '') + ': ' + corto(c ? c.innerText : '')); });
Array.prototype.slice.call(document.querySelectorAll('input')).filter(visible).slice(0,12).forEach(function(e){
  var c = e.closest('mat-form-field') || e.parentElement;
  r.casillas.push((e.type||'text') + ' [' + corto([e.getAttribute('placeholder'), e.getAttribute('data-placeholder'), e.getAttribute('aria-label'), e.getAttribute('formcontrolname')].filter(Boolean).join(' | ')) + ']: ' + corto(c ? c.innerText : '')); });
Array.prototype.slice.call(document.querySelectorAll('mat-option, [role="option"]')).filter(visible).slice(0,25).forEach(function(e){ r.opciones.push(corto(e.innerText)); });
var capa = document.querySelector('.cdk-overlay-container'); r.capas = capa ? capa.children.length : 0;
var p = document.querySelector('.mat-paginator-range-label, .mat-mdc-paginator-range-label'); r.paginador = p ? corto(p.innerText) : '';
return r;
"""


def _navegador(carpeta, visible):
    from selenium import webdriver

    usar_edge = os.environ.get("NAVEGADOR", "edge" if os.name == "nt" else "chrome").lower() == "edge"
    if usar_edge:
        from selenium.webdriver.edge.options import Options
    else:
        from selenium.webdriver.chrome.options import Options
    op = Options()
    op.add_experimental_option("prefs", {
        "download.default_directory": str(carpeta), "download.prompt_for_download": False,
        "download.directory_upgrade": True, "safebrowsing.enabled": True, "intl.accept_languages": "es-PE,es",
    })
    if not visible:
        op.add_argument("--headless=new")
        op.add_argument("--disable-gpu")
    for a in ("--window-size=1920,1080", "--disable-popup-blocking", "--log-level=3", "--lang=es-PE", "--no-sandbox", "--disable-dev-shm-usage"):
        op.add_argument(a)
    driver = webdriver.Edge(options=op) if usar_edge else webdriver.Chrome(options=op)
    try:   # en modo oculto hay que autorizar las descargas de forma explícita
        driver.execute_cdp_cmd("Page.setDownloadBehavior", {"behavior": "allow", "downloadPath": str(carpeta)})
    except Exception:
        pass
    return driver


def _visibles(driver, xpath):
    from selenium.webdriver.common.by import By
    salida = []
    for e in driver.find_elements(By.XPATH, xpath):
        try:
            if e.is_displayed():
                salida.append(e)
        except Exception:   # el elemento se redibujó mientras se revisaba
            pass
    return salida


def _esperar_visible(driver, xpaths, segundos, que):
    fin = time.time() + segundos
    while time.time() < fin:
        for xp in xpaths:
            v = _visibles(driver, xp)
            if v:
                return v[0]
        time.sleep(0.5)
    raise Detener(f"No apareció en pantalla: {que}.")


def _clic(driver, el):
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
    time.sleep(0.3)
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].click();", el)


def _control(driver, etiqueta, tipo, segundos=25):
    """Encuentra la lista ('lista') o la casilla ('casilla') que lleva esa etiqueta."""
    fin = time.time() + segundos
    while time.time() < fin:
        try:
            el = driver.execute_script(JS_CONTROL, etiqueta, tipo)
        except Exception:
            el = None
        if el is not None:
            return el
        time.sleep(0.5)
    raise Detener(f"No apareció en pantalla: {'la lista' if tipo == 'lista' else 'el campo'} «{etiqueta}».")


def _cerrar_lista(driver):
    from selenium.webdriver.common.action_chains import ActionChains
    from selenium.webdriver.common.keys import Keys
    try:
        ActionChains(driver).send_keys(Keys.ESCAPE).perform()
    except Exception:
        pass
    time.sleep(0.4)


def _abrir_lista(driver, control, forma):
    """Tres formas de desplegar la lista; se van alternando hasta que una funcione."""
    from selenium.webdriver.common.action_chains import ActionChains
    from selenium.webdriver.common.keys import Keys
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", control)
    time.sleep(0.3)
    try:
        if forma == 0:
            control.click()
        elif forma == 1:
            driver.execute_script("var t = arguments[0].querySelector('.mat-select-trigger, .mat-mdc-select-trigger') || arguments[0]; t.click();", control)
        else:
            driver.execute_script("arguments[0].focus();", control)
            ActionChains(driver).send_keys(Keys.ARROW_DOWN).perform()
    except Exception:
        pass


def _elegir(driver, etiqueta, opcion, segundos=90):
    """
    Abre la lista desplegable de esa etiqueta y elige la opción. Las opciones se
    cargan desde el servidor y pueden tardar: se reintenta hasta que aparezcan.
    """
    quiero = _norm(opcion)
    fin = time.time() + segundos
    vistas, abrio, vuelta = [], False, 0
    while time.time() < fin:
        control = _control(driver, etiqueta, "lista")
        _abrir_lista(driver, control, vuelta % 3)
        vuelta += 1
        opciones = []
        tope = time.time() + 6
        while time.time() < tope:
            try:
                opciones = driver.execute_script(JS_OPCIONES) or []
            except Exception:
                opciones = []
            if any(_norm(t) == quiero for _, t in opciones):
                break
            time.sleep(0.5)
        if opciones:
            abrio, vistas = True, [t for _, t in opciones]
        elegida = next((el for el, t in opciones if _norm(t) == quiero), None)
        if elegida is None:
            _cerrar_lista(driver)
            time.sleep(2)
            continue
        _clic(driver, elegida)
        time.sleep(1.0)
        _cerrar_lista(driver)
        try:   # se comprueba que la lista quedó con la opción puesta
            caja = driver.execute_script("var c = arguments[0].closest('mat-form-field') || arguments[0].parentElement; return c ? c.innerText : '';", _control(driver, etiqueta, "lista"))
        except Exception:
            caja = ""
        if quiero in _norm(caja):
            log(f"   {etiqueta}: {opcion}")
            return
        time.sleep(1)
    if not abrio:
        raise SinOpciones(f"La lista «{etiqueta}» no llegó a mostrar opciones.")
    raise Detener(f"La lista «{etiqueta}» no tiene la opción «{opcion}». Opciones vistas: {vistas[:20]}")


def _escribir_fecha(driver, etiqueta, fecha, invertido):
    """
    El campo muestra dd/mm/aaaa pero interpreta lo tecleado como mm/dd/aaaa
    (lo mismo que ya se vio en las automatizaciones anteriores), así que con
    invertido=True se teclea con día y mes cambiados. Quien decide si quedó
    bien no es este campo sino las fechas del Excel descargado: si no
    corresponden, el programa repite la descarga tecleando en el otro orden.
    """
    from selenium.webdriver.common.keys import Keys
    try:
        campo = _control(driver, etiqueta, "casilla", 20)
    except Detener:
        campo = _esperar_visible(driver, [
            f"//input[contains(@aria-label,'{etiqueta}') or contains(@placeholder,'{etiqueta}') or contains(@data-placeholder,'{etiqueta}')]",
        ], 10, f"el campo «{etiqueta}»")
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", campo)
    try:
        campo.click()
    except Exception:
        driver.execute_script("arguments[0].focus();", campo)
    time.sleep(0.2)
    campo.send_keys(Keys.CONTROL, "a")
    campo.send_keys(Keys.DELETE)
    time.sleep(0.15)
    for ch in fecha.strftime("%m/%d/%Y" if invertido else "%d/%m/%Y"):
        campo.send_keys(ch)
        time.sleep(0.05)
    time.sleep(0.2)
    driver.execute_script("document.activeElement.blur();")
    time.sleep(0.4)
    log(f"   {etiqueta}: se pidió {fecha.strftime('%d/%m/%Y')} y el campo muestra «{campo.get_attribute('value')}»")


def _filas_tabla(driver):
    try:
        n = len([f for f in _visibles(driver, "//table//tbody//tr | //mat-row") if f.text.strip()])
        if n:
            return n
        # segunda señal: el contador del pie de tabla deja de decir "0 of 0"
        pie = driver.execute_script("var p = document.querySelector('.mat-paginator-range-label, .mat-mdc-paginator-range-label'); return p ? p.innerText : '';") or ""
        return 1 if re.search(r"(?:of|de)\s+[1-9]", pie) else 0
    except Exception:   # la tabla se está redibujando
        return 0


def _esperar_resultados(driver, segundos):
    """El reporte no muestra ningún aviso de 'cargando': se espera a que aparezcan filas en la tabla."""
    fin = time.time() + segundos
    while time.time() < fin:
        n = _filas_tabla(driver)
        if n:
            for _ in range(10):   # se da por cargada cuando la cantidad de filas deja de cambiar
                time.sleep(2)
                m = _filas_tabla(driver)
                if m == n:
                    return True
                n = m
            return True
        time.sleep(2)
    return False


def _diagnostico(driver):
    """Pistas para corregir un fallo, sin datos del reporte (el registro de GitHub es público)."""
    try:
        log(f"   Pantalla al fallar: {driver.current_url.split('#')[-1]} | título «{driver.title}»")
        p = driver.execute_script(JS_PISTAS) or {}
        log(f"   Listas visibles: {p.get('listas')}")
        log(f"   Casillas visibles: {p.get('casillas')}")
        log(f"   Opciones desplegadas: {p.get('opciones')} | capas abiertas: {p.get('capas')} | paginador: «{p.get('paginador')}»")
        botones = sorted({b.text.strip() for b in _visibles(driver, "//button") if b.text.strip()})
        log(f"   Botones de acción: {[b for b in botones if any(x in b for x in ('Buscar', 'Exportar', 'Acceder'))]} (de {len(botones)} botones)")
        carpeta = os.environ.get("DIAGNOSTICO_DIR")   # solo para pruebas en tu PC; en GitHub no se define
        if carpeta:
            Path(carpeta).mkdir(parents=True, exist_ok=True)
            driver.save_screenshot(str(Path(carpeta) / "fallo.png"))
            (Path(carpeta) / "fallo.html").write_text(driver.page_source, encoding="utf-8")
            log(f"   Captura y HTML guardados en {carpeta}")
    except Exception:
        pass


def _abrir_reporte(driver, por_menu):
    boton = ["//button[contains(.,'Exportar a Excel')]"]
    if not por_menu:
        driver.get(URL_REPORTE)
        try:
            _esperar_visible(driver, boton, 30, "el reporte")
            return
        except Detener:
            log("   La dirección directa no abrió el reporte; se entra por el menú.")

    def texto(t):
        return [f"//*[normalize-space(text())='{t}']", f"//a[normalize-space(.)='{t}'] | //span[normalize-space(.)='{t}']"]

    for i, paso in enumerate(MENU):
        ultimo = i == len(MENU) - 1
        if not ultimo:   # un nivel ya desplegado no se vuelve a pulsar (se cerraría)
            siguiente = [e for xp in texto(MENU[i + 1]) for e in _visibles(driver, xp)]
            if siguiente:
                continue
        _clic(driver, _esperar_visible(driver, texto(paso), 25, f"el menú «{paso}»"))
        time.sleep(1.5)
    _esperar_visible(driver, boton, 40, "el reporte Recepción de Materia Prima")


def _sesion(driver, usuario, clave, inicio, fin, carpeta, invertido, intentos, por_menu):
    log("Abriendo la intranet e iniciando sesión...")
    try:
        driver.get(URL_LOGIN)
        campo_u = _esperar_visible(driver, ["//input[contains(@placeholder,'Usuario')]", "//input[@type='text' or @type='email' or not(@type)]"], 40, "el campo Usuario")
    except Exception:
        raise Detener("La intranet no abrió desde aquí (no apareció la pantalla de ingreso). "
                      "Si esto pasa en GitHub y en tu PC sí abre, el portal está bloqueando las conexiones de fuera.")
    campo_u.click()
    campo_u.send_keys(usuario)
    campo_c = _esperar_visible(driver, ["//input[@type='password']", "//input[contains(@placeholder,'Contraseña')]"], 15, "el campo Contraseña")
    campo_c.click()
    campo_c.send_keys(clave)
    _clic(driver, _esperar_visible(driver, ["//button[contains(.,'Acceder')]"], 15, "el botón Acceder"))
    _esperar_visible(driver, ["//*[contains(text(),'Aplicaciones Danper')]"], 60, "el panel principal (¿usuario o contraseña incorrectos?)")
    log("Sesión iniciada.")
    time.sleep(2)

    log("Abriendo el reporte Recepción de Materia Prima" + (" por el menú..." if por_menu else "..."))
    _abrir_reporte(driver, por_menu)
    time.sleep(3)

    log("Poniendo los filtros...")
    _elegir(driver, "Tipo Cultivo", TIPO_CULTIVO)
    time.sleep(1.5)   # Materia Prima se llena según el cultivo elegido
    _elegir(driver, "Materia Prima", MATERIA_PRIMA)
    _escribir_fecha(driver, "Fecha Inicio", inicio, invertido)
    _escribir_fecha(driver, "Fecha Fin", fin, invertido)

    hay = False
    for intento in range(1, intentos + 1):
        _clic(driver, _esperar_visible(driver, ["//button[contains(.,'Buscar')]"], 15, "el botón Buscar"))
        log(f"Buscar pulsado (intento {intento} de {intentos}); esperando la data, puede tardar...")
        if _esperar_resultados(driver, 150):
            hay = True
            break
    if not hay:
        log("La búsqueda no devolvió filas para esas fechas.")
        _diagnostico(driver)
        return None

    antes = set(os.listdir(carpeta))
    _clic(driver, _esperar_visible(driver, ["//button[contains(.,'Exportar a Excel')]"], 15, "el botón Exportar a Excel"))
    log("Exportando a Excel...")
    fin_espera = time.time() + 180
    while time.time() < fin_espera:
        nuevos = [f for f in set(os.listdir(carpeta)) - antes if not f.endswith((".crdownload", ".tmp"))]
        if nuevos:
            time.sleep(2)
            ruta = carpeta / sorted(nuevos, key=lambda f: (carpeta / f).stat().st_mtime)[-1]
            log(f"Excel descargado ({ruta.stat().st_size // 1024} KB).")
            return ruta
        time.sleep(0.5)
    raise Detener("Se pulsó Exportar a Excel pero el archivo no llegó a descargarse.")


def descargar(inicio, fin, carpeta, visible, invertido=True, intentos=3):
    usuario = os.environ.get("INTRANET_USUARIO", "").strip()
    clave = os.environ.get("INTRANET_CLAVE", "").strip() or usuario
    if not usuario:
        raise Detener("Falta el secreto INTRANET_USUARIO.")
    carpeta.mkdir(parents=True, exist_ok=True)

    # Primero se abre el reporte por su dirección directa. Si así las listas no cargan
    # sus opciones, se repite entrando por el menú, como se hace a mano.
    for por_menu in (False, True):
        for viejo in carpeta.glob("*.xls*"):
            viejo.unlink()
        try:
            driver = _navegador(carpeta, visible)
        except Exception as e:
            raise Detener(f"No se pudo abrir el navegador ({type(e).__name__}).")
        try:
            return _sesion(driver, usuario, clave, inicio, fin, carpeta, invertido, intentos, por_menu)
        except SinOpciones as e:
            _diagnostico(driver)
            if por_menu:
                raise
            log(f"   {e} Se repite entrando por el menú.")
        except Detener:
            _diagnostico(driver)
            raise
        except Exception as e:
            _diagnostico(driver)
            raise Detener(f"Fallo del navegador ({type(e).__name__}). Revisa las pistas de arriba.")
        finally:
            driver.quit()


# ============================================================
# 4. PROGRAMA PRINCIPAL
# ============================================================

def main():
    ap = argparse.ArgumentParser(description="Descarga Recepción de Materia Prima y actualiza datos.json")
    ap.add_argument("--dias", type=int, default=int(os.environ.get("DIAS_ATRAS") or 1), help="cuántos días hacia atrás traer (1 = solo ayer)")
    ap.add_argument("--excel", help="usar este Excel en vez de descargarlo de la intranet")
    ap.add_argument("--visible", action="store_true", help="mostrar el navegador (para pruebas en tu PC)")
    ap.add_argument("--solo-descarga", action="store_true", help="descargar y resumir, sin tocar datos.json")
    ap.add_argument("--hoy", help="fecha de hoy AAAA-MM-DD (solo para pruebas)")
    a = ap.parse_args()

    hoy = datetime.strptime(a.hoy, "%Y-%m-%d").date() if a.hoy else datetime.now(LIMA).date()
    dias = max(1, min(a.dias, 31))
    inicio, fin = hoy - timedelta(days=dias), hoy
    log(f"Hoy es {hoy.strftime('%d/%m/%Y')} (hora de Lima). Fecha Inicio {inicio.strftime('%d/%m/%Y')}, Fecha Fin {fin.strftime('%d/%m/%Y')}.")

    # Comprobación de fechas: lo descargado tiene que corresponder al rango pedido.
    d_ini, d_fin = (inicio - timedelta(days=1)).strftime("%Y-%m-%d"), (fin + timedelta(days=1)).strftime("%Y-%m-%d")

    def fechas_bien(fs):
        return sum(1 for f in fs if not (d_ini <= dia_de(f["r"]) <= d_fin)) <= 0.2 * len(fs)

    def rango(fs):
        vistos = sorted({dia_de(f["r"]) for f in fs})
        return f"{vistos[0]} a {vistos[-1]}"

    filas, mal = None, None
    if a.excel:
        filas = leer_excel(Path(a.excel))
        if filas and not fechas_bien(filas):
            mal, filas = rango(filas), None
    else:
        carpeta = Path(os.environ.get("DESCARGAS_DIR", "descargas")).resolve()
        # Primero se teclea la fecha como en las automatizaciones anteriores (día y mes cambiados).
        # Si el Excel no trae las fechas pedidas, se repite tecleando en el otro orden.
        for vuelta, invertido in enumerate((True, False)):
            if vuelta:
                log("Se repite la descarga escribiendo las fechas en el otro orden...")
            ruta = descargar(inicio, fin, carpeta, a.visible, invertido, 3 if vuelta == 0 else 2)
            leidas = leer_excel(ruta) if ruta else []
            if not leidas:
                continue
            if fechas_bien(leidas):
                filas, mal = leidas, None
                break
            mal = rango(leidas)
            log(f"El Excel trajo fechas que no se pidieron ({mal}): no se usa.")
    if mal:
        raise Detener(f"Las fechas del Excel no corresponden a lo pedido ({mal}). "
                      "El portal invirtió día y mes. No se publica nada.")
    if not filas:
        log("Sin datos nuevos: no se cambia nada.")
        return 0
    log(f"El Excel trae {len(filas)} recojos.")

    # Solo se publican días completos: desde Fecha Inicio hasta ayer. Lo de hoy todavía está a medias.
    objetivo = {(inicio + timedelta(days=i)).strftime("%Y-%m-%d") for i in range(dias)}
    por_dia = {}
    for f in filas:
        d = dia_de(f["r"])
        if d in objetivo:
            por_dia.setdefault(d, []).append(f)
    for d in sorted(por_dia):
        log(f"   {d[8:10]}/{d[5:7]}/{d[0:4]}: {len(por_dia[d])} recojos")
    if not por_dia:
        log("El Excel no trae recojos de los días pedidos: no se cambia nada.")
        return 0
    if a.solo_descarga:
        log("Prueba de descarga terminada (no se tocó datos.json).")
        return 0

    clave_admin = os.environ.get("ADMIN_CLAVE", "")
    if not clave_admin:
        raise Detener("Falta el secreto ADMIN_CLAVE.")
    contenido, dnis, it = abrir_datos(clave_admin)
    total = combinar(contenido, por_dia, "Automático (intranet)")
    guardar_datos(contenido, dnis, clave_admin, it, hoy)
    log(f"datos.json actualizado: {total} recojos en {len(por_dia)} día(s); {len(contenido['index']['dias'])} días en total; {len(dnis)} DNIs con acceso.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Detener as e:
        log(f"DETENIDO: {e}")
        sys.exit(1)
