"""
Prepara la plantilla de la Vista PPT a partir de TU PowerPoint del reporte diario.

Uso (cuando cambies el diseño del PPT):
    python herramientas/preparar_plantilla.py  Reporte_ESPARRAGO.pptx

Genera en lib/:
  - plantilla_reporte.pptx   tu PPT con cada valor reemplazado por una marca {{#Fnn}} (se llena en la página)
  - plantilla_campos.json    dónde va cada valor (posición, tamaño de letra, color…) para la vista previa
  - fondos/sNN.jpg           cada diapositiva sin esos valores (fondo de la vista previa y del PDF)
Necesita python-pptx y LibreOffice (soffice) + pdftoppm para los fondos.
"""
import copy, json, re, shutil, subprocess, sys, tempfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
EMU = 914400

# ----------------------------------------------------------------- qué valor va en cada cuadro
# slide -> { id de forma: [texto por párrafo] }   ''  = queda en blanco (todavía sin fuente de datos)
V = ""
MAPA = {
    1: {3: ["REPORTE DIARIO ESPÁRRAGO ({{fecha}})", "AÑO {{anio}}"]},
    2: {2: ["AVANCE DE CAMPAÑA ESPÁRRAGO - {{anio}}"], 4: ["CORTE HASTA SEM. {{sem}}"]},
    7: {2: ["INDICADORES DE\nLA\nSEMANA {{sem}}"]},
    # indicadores de producción (por ahora se escriben en la página): eg = Eficiencia Global, ef = Eficiencia,
    # mo = Mano de Obra – Producción, ap = Aprovechamiento, re = Rendimiento · acu = avance de campaña, sem = semana
    3: {52: ["{{p.mo.acu}}"], 56: ["{{p.mo.acu.dif}}"], 69: ["{{p.mo.sem}}"], 71: ["{{p.mo.sem.dif}}"],
        29: ["{{p.eg.acu.dif}}"], 39: ["{{p.eg.sem.dif}}"], 7: ["{{p.ap.acu}}"], 25: ["{{p.ap.acu.dif}}"],
        60: ["{{p.ap.sem}}"], 81: ["{{p.ap.sem.dif}}"], 83: ["{{p.re.sem}}"], 87: ["{{p.re.sem.dif}}"],
        91: ["{{p.re.acu}}"], 97: ["{{p.re.acu.dif}}"], 99: ["{{p.eg.acu}}"], 101: ["{{p.ef.acu}}"],
        105: ["{{p.ef.acu.dif}}"], 107: ["{{p.ef.sem}}"], 113: ["{{p.ef.sem.dif}}"], 115: ["{{p.eg.sem}}"]},
    4: {28: ["{{rec.proy}}"], 20: ["{{rec.eje}}"], 47: ["{{rec.sem}}"], 80: ["{{rec.avance}}"],
        77: ["{{dren.proy}}"], 50: ["{{dren.eje}}"], 46: ["{{dren.avance}}"], 33: ["{{dren.sem}}"], 122: ["{{obs.acu}}"], 123: ["{{obs.acu.dif}}"], 118: ["{{obs.sem}}"], 119: ["{{obs.sem.dif}}"]},
    5: {49: [V], 39: [V], 31: [V], 9: [V]},
    6: {47: ["{{mo.acu}}"], 51: ["{{mo.acu.dif}}"], 40: ["{{mo.sem.dif}}"], 22: ["{{mo.sem}}"], 25: ["{{mo.pers}}"], 52: ["{{recl.pct}}"], 56: ["{{recl.dif}}"], 16: ["{{recl.pen}}"], 20: ["{{recl.venta}}"],
        49: ["Objetivo: {{mo.acu.obj}}"], 39: ["Objetivo: {{mo.sem.obj}}"], 54: ["Objetivo: {{recl.obj}}"]},
    8: {15: ["{{rec.sem}}"], 30: [V], 19: ["KILOS RECEPCIONADOS – Sem. {{sem}}"], 21: ["KILOS EN STOCK – {{ddmm}}"]},
    12: {14: [V], 13: ["INSPECTORES EN EL DÍA – {{dia}} {{ddmm}}"]},
}
ZONAS = {9: "compositan", 10: "olmos", 11: "terceros"}
for sl, z in ZONAS.items():
    MAPA[sl] = {100: ["Prom: {{f.%s.elasmo.prom}}" % z], 102: ["Max: {{f.%s.elasmo.max}}" % z],
                104: ["Prom: {{f.%s.posturas.prom}}" % z], 106: ["Max: {{f.%s.posturas.max}}" % z],
                125: ["Prom: {{f.%s.larva.prom}}" % z], 127: ["Max: {{f.%s.larva.max}}" % z],
                129: ["Prom: {{f.%s.prodi.prom}}" % z], 131: ["Max: {{f.%s.prodi.max}}" % z]}
# calidades: (acumulado, ejecutado) de Small, Long>30 y C- y Ramas — los ids cambian en cada diapositiva
CAL = {9: [(147, 9), (17, 21), (28, 34)], 10: [(6, 30), (35, 41), (45, 49)], 11: [(6, 32), (37, 42), (46, 50)]}
for sl, pares in CAL.items():
    z = ZONAS[sl]
    for (acu, eje), k in zip(pares, ("sm", "lo", "cr")):
        MAPA[sl][acu] = ["{{c.%s.%s.acu}}" % (z, k)]
        MAPA[sl][eje] = ["{{c.%s.%s.eje}}" % (z, k)]
# líneas por categoría de cada caja fitosanitaria: se borran y se pone un solo cuadro con todas las líneas
LINEAS = {9: [29, 26, 56, 58, 67, 69, 92, 94, 96], 10: [97, 76, 86, 56, 58, 61, 67, 69, 74, 92], 11: [29, 97, 26, 31, 76, 86, 56, 58, 61, 67, 69, 74, 92, 94, 96]}
# 4 cajas del mismo ancho: título arriba, barra apilada a la izquierda, líneas al costado y Prom./Máx. a la derecha
CAJA_Y, CAJA_W, CAJA_H = 1.72, 3.02, 1.80
CAJA_X = [0.43, 3.58, 6.73, 9.88]
INDS = ["elasmo", "posturas", "larva", "prodi"]
TITULOS = ["SEMAFORIZACIÓN ELASMOSPALPUS", "POSTURAS", "LEPIDOPTERA", "PRODIPLOSIS"]
CAJAS = [(ind, x + 0.60, CAJA_Y + 0.42, 1.40, 1.30) for ind, x in zip(INDS, CAJA_X)]
BARRAS = [(ind, x + 0.15, CAJA_Y + 0.44, 0.40, 1.24) for ind, x in zip(INDS, CAJA_X)]
CAJAS_QUITAR = {9: [25, 36, 63, 88, 27, 40, 65, 90], 10: None, 11: None}     # recuadros y títulos viejos (se dibujan de nuevo)
PROM_MAX = [(100, 102), (104, 106), (125, 127), (129, 131)]
# imágenes con datos viejos (pantallazos de Power BI): se quitan
QUITAR = {3: [119], 6: [7], 8: [26], 12: [8]}
# donde estaba el pantallazo va un recuadro punteado gris "Imagen N" (ahí irá una imagen)
IMAGENES = {3: (119, "Imagen 1"), 6: (7, "Imagen 2"), 12: (8, "Imagen 3")}
# nota al pie de la diapositiva 3: se parte en dos líneas y se sube para que no se salga
PIE = {3: (40, 6.44, 0.52, 1300)}
# textos fijos que traen la semana o la fecha del reporte del 12/10
REEMPLAZOS = [(r"(?i)(semana de |semana |sem\. ?|sem ?|s\.)41", r"\g<1>{{sem}}"), (r"(?i)(semana |semana de )40", r"\g<1>{{sem_ant}}"),
              (r"Objetivo: 3\.50%", "Objetivo: {{obs.obj}}"),
              (r"^Objetivo: 83,0%$", "Objetivo: {{p.eg.obj}}"), (r"^Objetivo: 84,0%$", "Objetivo: {{p.ef.obj}}"),
              (r"^Objetivo: 0\.282$", "Objetivo: {{p.mo.obj}}"), (r"^Objetivo: 70 %$", "Objetivo: {{p.ap.obj}}"),
              (r"^Objetivo: 53\.5 %$", "Objetivo: {{p.re.obj}}"),
              (r"Información del mes de Julio - Agosto", "Información de {{recl.meses}}"),
              (r"12/10/2026", "{{fecha}}"), (r"12/10", "{{ddmm}}"), (r"Lun\.", "{{dia}}")]


def parrafos(el):
    """Párrafos del texto de una forma (p:txBody en formas, a:txBody en tablas)."""
    return etree.ElementBase.xpath(el, ".//p:txBody/a:p | .//a:txBody/a:p", namespaces=NS) if hasattr(etree, "ElementBase") else el.xpath(".//p:txBody/a:p")


def todas(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from todas(sh.shapes)


def abs_xfrm(sh):
    """Posición absoluta (pulgadas) aplicando las transformaciones de los grupos."""
    x, y, w, h = sh.left, sh.top, sh.width, sh.height
    el = sh._element.getparent()
    while el is not None and el.tag == P + "grpSp":
        xf = el.find(P + "grpSpPr/" + A + "xfrm")
        off, ext, choff, chext = (xf.find(A + t) for t in ("off", "ext", "chOff", "chExt"))
        sx = int(ext.get("cx")) / max(1, int(chext.get("cx")))
        sy = int(ext.get("cy")) / max(1, int(chext.get("cy")))
        x = int(off.get("x")) + (x - int(choff.get("x"))) * sx
        y = int(off.get("y")) + (y - int(choff.get("y"))) * sy
        w, h = w * sx, h * sy
        el = el.getparent()
    return [round(v / EMU, 4) for v in (x, y, w, h)]


def poner_texto(parrafo, texto):
    """Deja el párrafo con un solo run (el formato del primero) y ese texto."""
    runs = parrafo.findall(A + "r")
    if not runs:
        r = etree.SubElement(parrafo, A + "r"); etree.SubElement(r, A + "t")
        end = parrafo.find(A + "endParaRPr")
        if end is not None:
            rpr = copy.deepcopy(end); rpr.tag = A + "rPr"; r.insert(0, rpr)
            parrafo.remove(end); parrafo.append(end)
        runs = [r]
    for r in runs[1:]:
        parrafo.remove(r)
    for f in parrafo.findall(A + "fld") + parrafo.findall(A + "br"):
        parrafo.remove(f)
    runs[0].find(A + "t").text = texto


def estilo(parrafo, tema):
    r = parrafo.find(A + "r")
    rpr = r.find(A + "rPr") if r is not None else None
    if rpr is None:
        rpr = parrafo.find(A + "endParaRPr")
    st = {"sz": 18, "b": False, "i": False, "u": False, "c": "000000", "font": tema["minor"]}
    if rpr is not None:
        if rpr.get("sz"): st["sz"] = int(rpr.get("sz")) / 100
        st["b"] = rpr.get("b") == "1"; st["i"] = rpr.get("i") == "1"; st["u"] = rpr.get("u") not in (None, "none")
        fill = rpr.find(A + "solidFill")
        if fill is not None: st["c"] = color(fill, tema)
        lat = rpr.find(A + "latin")
        if lat is not None and lat.get("typeface"):
            tf = lat.get("typeface")
            st["font"] = tema["major"] if tf == "+mj-lt" else tema["minor"] if tf == "+mn-lt" else tf
    ppr = parrafo.find(A + "pPr")
    st["al"] = {"ctr": "center", "r": "right", "just": "justify"}.get(ppr.get("algn") if ppr is not None else None, "left")
    return st


def color(fill, tema):
    s = fill.find(A + "srgbClr")
    base, el = None, None
    if s is not None: base, el = s.get("val"), s
    s2 = fill.find(A + "schemeClr")
    if s2 is not None:
        n = {"tx1": "dk1", "tx2": "dk2", "bg1": "lt1", "bg2": "lt2"}.get(s2.get("val"), s2.get("val"))
        base, el = tema["colores"].get(n, "000000"), s2
    if base is None: return "000000"
    rgb = [int(base[i:i + 2], 16) for i in (0, 2, 4)]
    mod = el.find(A + "lumMod"); off = el.find(A + "lumOff")
    if mod is not None or off is not None:
        m = int(mod.get("val")) / 1e5 if mod is not None else 1; o = int(off.get("val")) / 1e5 if off is not None else 0
        rgb = [min(255, round(c * m + 255 * o)) for c in rgb]
    return "%02X%02X%02X" % tuple(rgb)


def leer_tema(prs):
    for rel in prs.slide_masters[0].part.rels.values():
        if rel.reltype.endswith("/theme"):
            x = etree.fromstring(rel.target_part.blob)
            cs = x.find(".//" + A + "clrScheme")
            col = {}
            for c in cs:
                v = c.find(A + "srgbClr"); s = c.find(A + "sysClr")
                col[etree.QName(c).localname] = v.get("val") if v is not None else (s.get("lastClr") if s is not None else "000000")
            fs = x.find(".//" + A + "fontScheme")
            return {"colores": col, "major": fs.find(A + "majorFont/" + A + "latin").get("typeface"), "minor": fs.find(A + "minorFont/" + A + "latin").get("typeface")}
    return {"colores": {}, "major": "Calibri", "minor": "Calibri"}


def body(sh):
    bp = sh._element.find(".//" + A + "bodyPr")  # a:bodyPr está dentro de p:txBody
    g = lambda k, d: int(bp.get(k)) / EMU if bp is not None and bp.get(k) is not None else d
    anchor = {"ctr": "middle", "b": "bottom"}.get(bp.get("anchor") if bp is not None else None, "top")
    return anchor, [g("lIns", 0.1), g("tIns", 0.05), g("rIns", 0.1), g("bIns", 0.05)]


def estilo_nota(parrafo):
    """Notas que empiezan con '*': negro, Arial 13, sin negrita ni cursiva (en el PPT)."""
    for r in parrafo.findall(A + "r"):
        rpr = r.find(A + "rPr")
        if rpr is None:
            rpr = etree.Element(A + "rPr"); r.insert(0, rpr)
        rpr.set("sz", "1300"); rpr.set("b", "0"); rpr.set("i", "0")
        if "u" in rpr.attrib: del rpr.attrib["u"]
        for t in list(rpr):
            if etree.QName(t).localname in ("solidFill", "gradFill", "noFill", "latin"):
                rpr.remove(t)
        fill = etree.Element(A + "solidFill"); etree.SubElement(fill, A + "srgbClr").set("val", "000000")
        rpr.insert(0, fill)
        pos = next((k for k, t in enumerate(rpr) if etree.QName(t).localname in ("ea", "cs", "sym", "hlinkClick", "rtl", "extLst")), len(rpr))
        etree.Element(A + "latin")
        lat = etree.Element(A + "latin"); lat.set("typeface", "Arial"); rpr.insert(pos, lat)


# Etiquetas de la diapositiva 8 que se centran en su recuadro: id → (x, ancho)
CENTRAR = {8: {23: (1.51, 3.73), 31: (8.36, 3.73), 15: (1.51, 3.73), 30: (8.36, 3.73)}}


def nuevo_cuadro(slide, modelo, x, y, w, h, texto, nid):
    el = copy.deepcopy(modelo._element)
    nv = el.find(".//" + P + "cNvPr"); nv.set("id", str(nid)); nv.set("name", "Lineas %d" % nid)
    xf = el.find(".//" + A + "xfrm")
    xf.find(A + "off").set("x", str(int(x * EMU))); xf.find(A + "off").set("y", str(int(y * EMU)))
    xf.find(A + "ext").set("cx", str(int(w * EMU))); xf.find(A + "ext").set("cy", str(int(h * EMU)))
    bp = el.find(".//" + A + "bodyPr"); bp.set("wrap", "none")
    for t in bp: bp.remove(t)
    ps = parrafos(el)
    for p_ in ps[1:]: p_.getparent().remove(p_)
    poner_texto(ps[0], texto)
    rpr = ps[0].find(A + "r/" + A + "rPr")
    if rpr is not None: rpr.set("sz", "1100")
    slide.shapes._spTree.append(el)
    return el


def recuadro_imagen(x, y, w, h, texto, nid):
    """Recuadro punteado gris claro con el texto al centro (marca dónde va una imagen)."""
    return etree.fromstring(
        '<p:sp xmlns:p="%s" xmlns:a="%s"><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 5000"/></a:avLst></a:prstGeom>'
        '<a:solidFill><a:srgbClr val="F4F4F4"/></a:solidFill><a:ln w="15875"><a:solidFill><a:srgbClr val="A6A6A6"/></a:solidFill><a:prstDash val="dash"/></a:ln></p:spPr>'
        '<p:txBody><a:bodyPr anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="es-PE" sz="1600"><a:solidFill><a:srgbClr val="A6A6A6"/></a:solidFill>'
        '<a:latin typeface="Arial"/></a:rPr><a:t>%s</a:t></a:r></a:p></p:txBody></p:sp>'
        % (NS["p"], NS["a"], nid, texto, x * EMU, y * EMU, w * EMU, h * EMU, texto))


def tabla_xml(x, y, cols, alto, filas, nid):
    """Tabla nativa (Zona | Kg | %) en lugar del pantallazo."""
    w = sum(cols)
    g = ['<p:graphicFrame xmlns:p="%s" xmlns:a="%s"><p:nvGraphicFramePr><p:cNvPr id="%d" name="Tabla zonas"/><p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>' % (NS["p"], NS["a"], nid),
         '<p:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></p:xfrm>' % (x * EMU, y * EMU, w * EMU, alto * len(filas) * EMU),
         '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table"><a:tbl><a:tblPr firstRow="1"/><a:tblGrid>']
    g += ['<a:gridCol w="%d"/>' % (c * EMU) for c in cols] + ["</a:tblGrid>"]
    borde = '<a:solidFill><a:srgbClr val="BFBFBF"/></a:solidFill>'
    for i, fila in enumerate(filas):
        g.append('<a:tr h="%d">' % (alto * EMU))
        for j, txt in enumerate(fila):
            b = ' b="1"' if i in (0, len(filas) - 1) else ""
            al = ' algn="r"' if j else ""
            fill = '<a:solidFill><a:srgbClr val="F2F2F2"/></a:solidFill>' if i == 0 else '<a:noFill/>'
            g.append('<a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:pPr%s/><a:r><a:rPr lang="es-PE" sz="1000"%s><a:solidFill><a:srgbClr val="262626"/></a:solidFill></a:rPr><a:t>%s</a:t></a:r><a:endParaRPr lang="es-PE" sz="1000"/></a:p></a:txBody>'
                     '<a:tcPr marL="60000" marR="60000" marT="20000" marB="20000" anchor="ctr">'
                     '<a:lnL w="6350">%s</a:lnL><a:lnR w="6350">%s</a:lnR><a:lnT w="6350">%s</a:lnT><a:lnB w="6350">%s</a:lnB>%s</a:tcPr></a:tc>' % (al, b, txt, borde, borde, borde, borde, fill))
        g.append("</a:tr>")
    g.append("</a:tbl></a:graphicData></a:graphic></p:graphicFrame>")
    return etree.fromstring("".join(g))


def main(origen, destino):
    prs = Presentation(origen)
    tema = leer_tema(prs)
    campos, n, barras, huecos = [], 0, [], []
    fondo = Presentation(origen)          # copia para los fondos (se le quitan los textos que cambian)

    def marcar(sl_i, sh, textos):
        """Cada párrafo pasa a ser un campo editable {{#Fnn}}. textos[k] = None → párrafo vacío que se deja tal cual."""
        nonlocal n
        ps = parrafos(sh._element)
        anchor, ins = body(sh)
        box = abs_xfrm(sh)
        bp = sh._element.find(".//" + A + "bodyPr")
        wrap = "none" if bp is not None and bp.get("wrap") == "none" else "square"
        k = [round(box[2] / max(1e-6, sh.width / EMU), 6), round(box[3] / max(1e-6, sh.height / EMU), 6)] if sh.width and sh.height else [1, 1]
        titulo = sh.is_placeholder and "TITLE" in str(sh.placeholder_format.type)
        cen = CENTRAR.get(sl_i, {}).get(sh.shape_id)
        if cen:
            box = [cen[0], box[1], cen[1], box[3]]
        reg = {"s": sl_i, "box": box, "anchor": anchor, "ins": [round(v, 4) for v in ins], "wrap": wrap, "k": k, "p": []}
        for j, p_ in enumerate(ps):
            if j >= len(textos):
                p_.getparent().remove(p_); continue
            if textos[j] is None:
                continue
            n += 1
            st = estilo(p_, tema)
            if titulo and st["font"] == tema["minor"]:
                st["font"] = tema["major"]
            st.update({"id": "F%d" % n, "t": textos[j]})
            poner_texto(p_, "{{#F%d}}" % n)
            if sl_i == 3 and textos[j].lstrip().startswith("Objetivo") and box[1] < 3.6:
                # objetivos de Eficiencia Global, Eficiencia y Mano de Obra: 17 pt
                st["sz"] = 17
                for r_ in p_.findall(A + "r"):
                    rpr = r_.find(A + "rPr")
                    if rpr is None:
                        rpr = etree.Element(A + "rPr"); r_.insert(0, rpr)
                    rpr.set("sz", "1700")
            if textos[j].lstrip().startswith("*"):
                st.update({"c": "000000", "b": False, "i": False, "u": False, "font": "Arial", "sz": 13})
                estilo_nota(p_)
            cen = CENTRAR.get(sl_i, {}).get(sh.shape_id)
            if cen:
                st["al"] = "center"
                ppr = p_.find(A + "pPr")
                if ppr is None:
                    ppr = etree.Element(A + "pPr"); p_.insert(0, ppr)
                ppr.set("algn", "ctr")
            reg["p"].append(st)
        if reg["p"]:
            campos.append(reg)

    for i, (sl, sf) in enumerate(zip(prs.slides, fondo.slides), 1):
        por_id = {sh.shape_id: sh for sh in todas(sl.shapes)}
        por_id_f = {sh.shape_id: sh for sh in todas(sf.shapes)}
        # 0) recuadros "Imagen N" en el lugar de los pantallazos (van en el fondo y en el PowerPoint)
        if i in IMAGENES:
            sid, txt = IMAGENES[i]
            x, y, w, h = abs_xfrm(por_id[sid])
            huecos.append({"s": i, "box": [x, y, w, h], "nombre": txt})
            for presen in (sl, sf):
                presen.shapes._spTree.append(recuadro_imagen(x, y, w, h, txt, 700 + i))
        # 0b) nota al pie más arriba, en dos líneas
        if i in PIE:
            sid, y, h, sz = PIE[i]
            for d in (por_id, por_id_f):
                el = d[sid]._element
                xf = el.find(".//" + A + "xfrm")
                xf.find(A + "off").set("y", str(int(y * EMU))); xf.find(A + "ext").set("cy", str(int(h * EMU)))
                el.find(".//" + A + "bodyPr").set("wrap", "square")
                for r_ in el.iter(A + "rPr"):
                    r_.set("sz", str(sz))
        # 1) quitar pantallazos y líneas de categorías
        for sid in QUITAR.get(i, []) + LINEAS.get(i, []):
            for d in (por_id, por_id_f):
                sh = d.get(sid)
                if sh is not None and sh._element.getparent() is not None:
                    sh._element.getparent().remove(sh._element)
        # 1b) cajas fitosanitarias: se quitan los recuadros y títulos viejos y se dibujan 4 iguales
        if i in ZONAS:
            rect_modelo = copy.deepcopy(por_id[36]._element)
            tit_modelo = copy.deepcopy(por_id[40]._element)
            viejos = [sid for sid, sh in por_id.items() if sh.has_text_frame and sh.text_frame.text.strip() in TITULOS] + \
                     [sid for sid, sh in por_id.items() if sh.name.startswith("Rectángulo: esquinas") and sh.top is not None and sh.top < 4 * EMU and sh.width > 2 * EMU] + \
                     [sid for sid, sh in por_id.items() if sh.name.startswith("Rectángulo: esquinas") and sh.top is not None and abs(sh.width / EMU - 3.12) < 0.02]
            for sid in set(viejos):
                for d in (por_id, por_id_f):
                    sh = d.get(sid)
                    if sh is not None and sh._element.getparent() is not None:
                        sh._element.getparent().remove(sh._element)
            for k, x in enumerate(CAJA_X):
                for presen in (sl, sf):
                    r_ = copy.deepcopy(rect_modelo); r_.find(".//" + P + "cNvPr").set("id", str(800 + k))
                    xf = r_.find(".//" + A + "xfrm")
                    xf.find(A + "off").set("x", str(int(x * EMU))); xf.find(A + "off").set("y", str(int(CAJA_Y * EMU)))
                    xf.find(A + "ext").set("cx", str(int(CAJA_W * EMU))); xf.find(A + "ext").set("cy", str(int(CAJA_H * EMU)))
                    presen.shapes._spTree.append(r_)
                t_ = copy.deepcopy(tit_modelo); t_.find(".//" + P + "cNvPr").set("id", str(810 + k))
                xf = t_.find(".//" + A + "xfrm")
                xf.find(A + "off").set("x", str(int((x + 0.05) * EMU))); xf.find(A + "off").set("y", str(int((CAJA_Y + 0.07) * EMU)))
                xf.find(A + "ext").set("cx", str(int((CAJA_W - 0.10) * EMU))); xf.find(A + "ext").set("cy", str(int(0.30 * EMU)))
                ps_ = parrafos(t_)
                for p_ in ps_[1:]: p_.getparent().remove(p_)
                poner_texto(ps_[0], TITULOS[k])
                ppr = ps_[0].find(A + "pPr")
                if ppr is None: ppr = etree.SubElement(ps_[0], A + "pPr"); ps_[0].insert(0, ppr)
                ppr.set("algn", "ctr")
                sl.shapes._spTree.append(t_)
                sf.shapes._spTree.append(copy.deepcopy(t_))
            # Prom./Máx. a la derecha de cada caja
            for k, (ip, im) in enumerate(PROM_MAX):
                for sid, dy in ((ip, 0.55), (im, 0.92)):
                    for d in (por_id, por_id_f):
                        xf = d[sid]._element.find(".//" + A + "xfrm")
                        xf.find(A + "off").set("x", str(int((CAJA_X[k] + 1.98) * EMU))); xf.find(A + "off").set("y", str(int((CAJA_Y + dy) * EMU)))
                        xf.find(A + "ext").set("cx", str(int(1.02 * EMU))); xf.find(A + "ext").set("cy", str(int(0.32 * EMU)))
                        for r_ in d[sid]._element.iter(A + "rPr"):
                            r_.set("sz", "1300")
            por_id = {sh.shape_id: sh for sh in todas(sl.shapes)}
            por_id_f = {sh.shape_id: sh for sh in todas(sf.shapes)}
            barras.extend({"s": i, "z": ZONAS[i], "ind": ind, "box": [x, y, w, h]} for ind, x, y, w, h in BARRAS)
        # 2) valores mapeados
        for sid, textos in MAPA.get(i, {}).items():
            if sid in por_id and por_id[sid].has_text_frame:
                marcar(i, por_id[sid], textos)
                for p_ in parrafos(por_id_f[sid]._element):
                    poner_texto(p_, "")
        # 3) TODOS los demás textos también pasan a ser campos (editables, se pueden mover y cambiar de letra);
        #    los que traen la semana o la fecha del reporte se actualizan solos
        for sid, sh in list(por_id.items()):
            if not sh.has_text_frame or sid in MAPA.get(i, {}) or sh._element.getparent() is None:
                continue
            ps = parrafos(sh._element)
            textos = ["".join(t.text or "" for t in p_.iter(A + "t")) for p_ in ps]
            nuevos = []
            for t in textos:
                if not t.strip():
                    nuevos.append(None); continue
                for pat, rep_ in REEMPLAZOS:
                    t = re.sub(pat, rep_, t)
                nuevos.append(t)
            if any(x is not None for x in nuevos):
                marcar(i, sh, nuevos)
                for p_ in parrafos(por_id_f[sid]._element):
                    poner_texto(p_, "")
        # 4) cuadros nuevos con las líneas de cada caja fitosanitaria
        if i in ZONAS:
            modelo = por_id[100]
            for ind, x, y, w, h in CAJAS:
                n += 1
                el = nuevo_cuadro(sl, modelo, x, y, w, h, "{{#F%d}}" % n, 900 + n)
                p_ = parrafos(el)[0]
                st = estilo(p_, tema); st.update({"id": "F%d" % n, "t": "{{f.%s.%s.lineas}}" % (ZONAS[i], ind), "sz": 11, "lineas": "f.%s.%s" % (ZONAS[i], ind)})
                campos.append({"s": i, "box": [x, y, w, h], "anchor": "top", "ins": [0.1, 0.05, 0.1, 0.05], "wrap": "none", "k": [1, 1], "p": [st]})
        # 5) tabla de zonas en lugar del pantallazo (diapositiva 8)
        if i == 8:
            x, y, cols, alto = 1.51, 3.62, [1.25, 1.5, 0.98], 0.27
            filas, regs = [["Zona", "Kg Recepcionados", "% Kg Recep"]], []
            for r in range(1, 5):
                fila = []
                for c in range(3):
                    n += 1; fila.append("{{#F%d}}" % n)
                    regs.append((r, c, "F%d" % n))
                filas.append(fila)
            sl.shapes._spTree.append(tabla_xml(x, y, cols, alto, filas, 990))
            sf.shapes._spTree.append(tabla_xml(x, y, cols, alto, [filas[0]] + [[" ", " ", " "]] * 4, 990))
            for r, c, fid in regs:
                cx = x + sum(cols[:c])
                ult = r == 4
                campos.append({"s": 8, "box": [cx, y + alto * r, cols[c], alto], "anchor": "middle", "ins": [0.066, 0.02, 0.066, 0.02], "wrap": "none", "k": [1, 1], "tabla": True,
                               "p": [{"id": fid, "t": "{{tz.%d.%d}}" % (r, c), "sz": 10, "b": ult, "i": False, "u": False, "c": "262626", "font": tema["minor"], "al": "right" if c else "left"}]})
    destino.mkdir(parents=True, exist_ok=True)
    prs.save(destino / "plantilla_reporte.pptx")
    (destino / "plantilla_campos.json").write_text(json.dumps({"w": prs.slide_width / EMU, "h": prs.slide_height / EMU, "n": len(prs.slides), "tema": tema, "campos": campos, "barras": barras, "imagenes": huecos,
                                                       "archivos": {str(k): sl.part.partname.lstrip("/") for k, sl in enumerate(prs.slides, 1)}}, ensure_ascii=False, indent=0), encoding="utf-8")

    # fondos: la copia sin los textos que cambian → PDF (LibreOffice) → JPG
    tmp = Path(tempfile.mkdtemp())
    fondo.save(tmp / "fondo.pptx")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(tmp / "fondo.pptx")], check=True, capture_output=True, timeout=300)
    (destino / "fondos").mkdir(exist_ok=True)
    subprocess.run(["pdftoppm", "-r", "150", "-jpeg", "-jpegopt", "quality=85", str(tmp / "fondo.pdf"), str(tmp / "s")], check=True)
    for k, f in enumerate(sorted(tmp.glob("s-*.jpg")), 1):
        shutil.copy(f, destino / "fondos" / ("s%02d.jpg" % k))
    print("Listo:", len(campos), "cuadros con valores en", len(prs.slides), "diapositivas →", destino)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("Uso: python herramientas/preparar_plantilla.py  Reporte_ESPARRAGO.pptx")
    main(sys.argv[1], Path(__file__).resolve().parent.parent / "lib")
